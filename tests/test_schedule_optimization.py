"""Synthetic regressions: no accounts, real events or participant information."""
import hashlib
import importlib.util
import itertools
import json
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'gamemaster-skill/scripts'))
import schedule_engine as engine


def match(mid, project='P', athletes=(), **kw):
    return dict(id=mid, project_id=project, athletes=list(athletes), predecessors=[], **kw)


def fixture(matches, start='10:30', end='12:30', courts=(1, 2), **kw):
    return dict(event_id='SYNTHETIC', courts=list(courts), slot_minutes=15,
                sessions=[dict(start='2030-01-01T' + start, end='2030-01-01T' + end, section=1)],
                rest_basis='elapsed', rest_minutes=15, conflict_scope='known', matches=matches,
                attempts=4, random_seed=7,
                constraint_sources={'*': {'kind': 'user', 'reference': 'Synthetic fixture specification'}}, **kw)


def candidate(data, placements):
    slots, matches, _, _ = engine.prepare(data)
    return dict(event_id=data['event_id'], slots=slots,
                input_sha256=hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
                assignments=[engine.record_for(data, matches[mid], slots[idx], court)
                             for mid, (idx, court) in placements.items()], unscheduled=[])


def locations(result):
    return {r['match_id']: r['slot_index'] for r in result['assignments']}


class SchedulingOptimization(unittest.TestCase):
    def test_final_moves_from_1145_to_1130(self):
        semi = match('semi', fixed={'slot_index': 0})
        final = match('final'); final['predecessors'] = ['semi']
        data = fixture([semi, final], start='11:00', focus_matches=['final'])
        before = candidate(data, {'semi': (0, 1), 'final': (3, 1)})
        audit = engine.acceptance(data, before)
        self.assertTrue(audit['hard_constraints']['complete'])
        self.assertFalse(audit['ready_for_export'])
        reason = audit['waiting_explanations'][0]
        self.assertEqual(reason['earliest_legal_start_with_other_matches_fixed'], '2030-01-01T11:30')
        self.assertEqual(reason['predecessors'][0]['end'], '2030-01-01T11:15')
        after = engine.improve(data, before)
        self.assertEqual(locations(after)['final'], 2)
        self.assertTrue(after['acceptance']['ready_for_export'])
        self.assertEqual(after['optimization']['before']['makespan'] - after['metrics']['makespan'], 15)

    def test_rest_begins_at_predecessor_end(self):
        final = match('final'); final['predecessors'] = ['semi']
        data = fixture([match('semi', fixed={'slot_index': 0}), final])
        early = candidate(data, {'semi': (0, 1), 'final': (1, 2)})
        self.assertFalse(engine.validate(data, early)['complete'])
        result = engine.schedule(data)
        self.assertEqual(locations(result)['final'], 2)
        self.assertEqual(next(r['start'] for r in result['assignments'] if r['match_id'] == 'final'), '2030-01-01T11:00')

    def test_priority_allows_equal_finish(self):
        data = fixture([match('priority', 'XD'), match('other', 'MS')], priority_projects=['XD'])
        before = candidate(data, {'priority': (0, 1), 'other': (1, 2)})
        result = engine.improve(data, before)
        self.assertEqual(locations(result), {'priority': 0, 'other': 0})
        self.assertEqual(locations(engine.schedule(data)), locations(result))

    def test_cross_project_rest_rejects_move(self):
        data = fixture([match('other', 'MS', ['athlete-A'], fixed={'slot_index': 1}),
                        match('final', 'XD', ['athlete-A'])], focus_matches=['final'])
        before = candidate(data, {'other': (1, 1), 'final': (3, 2)})
        ctx = engine.prepare(data); placed = {r['match_id']: r for r in before['assignments']}
        self.assertIsNone(engine.moved_records(data, ctx, placed, {'final': 2}))
        self.assertEqual(locations(engine.improve(data, before))['final'], 3)

    def test_fixed_successor_of_cross_project_match_rejects_move(self):
        data = fixture([match('fixed', 'MS', ['athlete-A'], fixed={'slot_index': 1}),
                        match('movable', 'XD', ['athlete-A'])])
        before = candidate(data, {'fixed': (1, 1), 'movable': (3, 2)})
        ctx = engine.prepare(data); placed = {r['match_id']: r for r in before['assignments']}
        self.assertIsNone(engine.moved_records(data, ctx, placed, {'movable': 0}))
        self.assertEqual(locations(engine.improve(data, before))['fixed'], 1)

    def test_chain_cannot_break_fixed_downstream_dependency(self):
        downstream = match('downstream', fixed={'slot_index': 4}); downstream['predecessors'] = ['parent']
        data = fixture([match('parent'), downstream, match('other', 'Q')])
        before = candidate(data, {'parent': (2, 1), 'downstream': (4, 1), 'other': (5, 2)})
        ctx = engine.prepare(data); placed = {r['match_id']: r for r in before['assignments']}
        # A swap advances the other match but delays its fixed successor's feeder.
        self.assertIsNone(engine.moved_records(data, ctx, placed, {'parent': 5, 'other': 2}))
        self.assertIsNone(engine.moved_records(data, ctx, placed, {'parent': 1, 'downstream': 3}))

    def test_round_cohesion_is_not_hard(self):
        data = fixture([match('one', athletes=['athlete-A'], stage=1, round=1),
                        match('two', athletes=['athlete-A'], stage=1, round=1)])
        result = engine.schedule(data)
        self.assertTrue(result['validation']['complete'])
        self.assertEqual(sorted(locations(result).values()), [0, 2])
        self.assertEqual(result['metrics']['round_splits'], 1)

    def test_final_bronze_opposite_paths_parallel(self):
        a = match('semi-A', possible_athletes=['a', 'b'], fixed={'slot_index': 0})
        b = match('semi-B', possible_athletes=['c', 'd'], fixed={'slot_index': 0})
        finals = []
        for mid, kind in [('final', 'winner'), ('bronze', 'loser')]:
            m = match(mid, possible_athletes=['a', 'b', 'c', 'd'],
                      sides=[dict(kind=kind, match_id=p) for p in ('semi-A', 'semi-B')])
            m['predecessors'] = ['semi-A', 'semi-B']; finals.append(m)
        data = fixture([a, b] + finals, outcome_disjoint_pairs=[['final', 'bronze']])
        data['conflict_scope'] = 'possible'
        result = engine.schedule(data)
        self.assertTrue(result['validation']['complete'])
        self.assertEqual(locations(result)['final'], 2)
        self.assertEqual(locations(result)['bronze'], 2)
        bad = deepcopy(data); bad['matches'][-1]['sides'][0]['kind'] = 'winner'
        with self.assertRaises(ValueError): engine.prepare(bad)

    def test_strict_finish_requires_user_source_and_is_hard(self):
        data = fixture([match('priority', 'XD'), match('other', 'MS')], priority_projects=['XD'],
                       strict_finish_order=[{'before': 'XD', 'after': 'MS'}])
        with self.assertRaises(ValueError): engine.prepare(data)
        data['strict_finish_order'][0]['source'] = {'kind': 'user', 'reference': 'Synthetic strict order'}
        result = engine.schedule(data)
        self.assertTrue(result['validation']['complete'])
        self.assertLess(locations(result)['priority'], locations(result)['other'])
        self.assertFalse(engine.validate(data, candidate(data, {'priority': (0, 1), 'other': (0, 2)}))['valid'])

    def test_round_group_move_survives_single_move_fragmentation(self):
        data = fixture([match('one', round=1), match('two', round=1),
                        match('end', fixed={'slot_index': 4})], courts=(1, 2, 3))
        before = candidate(data, {'one': (2, 2), 'two': (2, 3), 'end': (4, 1)})
        result = engine.improve(data, before)
        self.assertEqual(result['metrics']['round_splits'], 0)
        self.assertEqual(locations(result)['one'], 0)
        self.assertTrue(result['acceptance']['quality_checked'])
        self.assertEqual(result['optimization']['moves'][0]['kind'], 'round_group')

    def test_upstream_chain_neighborhood_available(self):
        child = match('child'); child['predecessors'] = ['parent']
        data = fixture([match('parent'), child])
        before = candidate(data, {'parent': (2, 1), 'child': (4, 1)})
        context = engine.prepare(data); placed = {r['match_id']: r for r in before['assignments']}
        chains = dict(engine.proposals(context, placed, engine.search_settings(data)))['ancestor_chain']
        targets = list(chains)
        self.assertIn({'parent': 0, 'child': 2}, targets)
        self.assertIsNotNone(engine.moved_records(data, context, placed, targets[-1]))

    def test_round_cohort_moves_without_moving_fixed_earlier_cohort(self):
        data = fixture([match('fixed-A', round=1, fixed={'slot_index': 0}),
                        match('fixed-B', round=1, fixed={'slot_index': 0}),
                        match('one', round=1), match('two', round=1),
                        match('end', fixed={'slot_index': 4})])
        before = candidate(data, {'fixed-A': (0, 1), 'fixed-B': (0, 2),
                                  'one': (2, 1), 'two': (2, 2), 'end': (4, 1)})
        result = engine.improve(data, before)
        self.assertEqual(locations(result)['one'], 1)
        self.assertEqual(locations(result)['two'], 1)
        self.assertEqual(result['optimization']['moves'][0]['kind'], 'round_group')

    def test_chain_improves_when_single_moves_are_blocked_or_worsen_rounds(self):
        child = match('child', round=2); child['predecessors'] = ['parent']
        data = fixture([match('parent', round=1), match('buddy', round=1, fixed={'slot_index': 2}), child])
        before = candidate(data, {'parent': (2, 1), 'buddy': (2, 2), 'child': (4, 1)})
        result = engine.improve(data, before)
        self.assertEqual(result['optimization']['moves'][0]['kind'], 'ancestor_chain')
        self.assertEqual(result['metrics']['makespan'], 45)

    def test_round_preference_explained_as_tradeoff(self):
        data = fixture([match('one', round=1), match('two', round=1, fixed={'slot_index': 2}),
                        match('end', fixed={'slot_index': 4})], focus_matches=['one'])
        result = engine.improve(data, candidate(data, {'one': (2, 1), 'two': (2, 2), 'end': (4, 1)}))
        self.assertEqual(locations(result)['one'], 2)
        report = result['acceptance']['waiting_explanations'][0]
        self.assertEqual(report['earliest_legal_start_with_other_matches_fixed'], '2030-01-01T10:30')
        reasons = report['earlier_scenes'][0]['alternatives'][0]['reasons']
        self.assertEqual(reasons[0]['category'], 'optimization_tradeoff')
        self.assertEqual(reasons[0]['objective'], 'round_splits')

    def test_court_rematching_preserves_specific_court(self):
        data = fixture([match('flex'), match('limited', allowed_courts=[1])])
        before = candidate(data, {'flex': (2, 2), 'limited': (2, 1)})
        placed = {r['match_id']: r for r in before['assignments']}
        trial = engine.moved_records(data, engine.prepare(data), placed, {'flex': 0, 'limited': 0})
        self.assertEqual(trial['limited']['court'], 1)
        self.assertEqual(trial['flex']['court'], 2)

    def test_window_blocker_does_not_blame_satisfied_constraints(self):
        data = fixture([match('one', not_before='2030-01-01T11:00', allowed_courts=[1])])
        result = engine.schedule(data)
        item = result['acceptance']['waiting_explanations'][0]
        codes = [r['constraint'] for r in item['earlier_scenes'][0]['alternatives'][0]['reasons']]
        self.assertEqual(codes, ['matches.one.not_before'])

    def test_swap_frees_early_slot(self):
        data = fixture([match('priority', 'XD'), match('other', 'MS')], courts=(1,), priority_projects=['XD'])
        before = candidate(data, {'other': (0, 1), 'priority': (1, 1)})
        result = engine.improve(data, before)
        self.assertEqual(locations(result)['priority'], 0)
        self.assertEqual(result['optimization']['moves'][0]['kind'], 'pair_swap')

    def test_metrics_are_distinct_and_duration_is_actual(self):
        data = fixture([match('one', duration_minutes=5), match('two')])
        result = candidate(data, {'one': (0, 1), 'two': (3, 1)})
        values = engine.metrics(data, result)
        self.assertEqual((values['total_matches'], values['occupied_scenes'], values['scene_span'], values['makespan']), (2, 2, 4, 60))
        self.assertEqual(engine.objective_order(data), list(engine.DEFAULT_OBJECTIVES))
        data['objective_order'] = list(reversed(engine.DEFAULT_OBJECTIVES))
        self.assertEqual(engine.objective_order(data)[0], 'start_times')

    def test_unknown_sources_block_formal_export_without_invention(self):
        data = fixture([match('one')]); del data['constraint_sources']
        result = engine.schedule(data)
        self.assertTrue(result['validation']['complete'])
        self.assertFalse(result['acceptance']['ready_for_export'])
        self.assertTrue(any(x['source']['kind'] == 'unverified_input' for x in result['constraint_sources']))

    def test_search_limit_is_not_proof_or_acceptance(self):
        data = fixture([match('one', fixed={'slot_index': 4})], local_search={'max_evaluations': 1})
        result = engine.schedule(data)
        self.assertFalse(result['acceptance']['quality_checked'])
        self.assertEqual(result['acceptance']['local_search']['status'], 'search_limit')
        self.assertFalse(result['optimality_proven'])

    def test_higher_objectives_never_worsen_and_tiny_exhaustive_oracle(self):
        data = fixture([match('a', 'Q'), match('b', 'P'), match('c', 'P')],
                       start='10:30', end='11:15', courts=(1,), priority_projects=['P'])
        data['rest_minutes'] = 0
        result = engine.schedule(data)
        scores = []
        for perm in itertools.permutations(range(3)):
            trial = candidate(data, {mid: (idx, 1) for mid, idx in zip(('a', 'b', 'c'), perm)})
            self.assertTrue(engine.validate(data, trial)['complete'])
            # Independent numerical oracle for this no-conflict fixture.
            ends = {'Q': 15 * (perm[0] + 1), 'P': 15 * (max(perm[1:]) + 1)}
            scores.append((45, (ends['P'],), sum(ends.values()), 0, 0, 45, tuple(15*x for x in perm)))
        self.assertEqual(engine.score(data, result['metrics']), min(scores))
        for move in result['optimization']['moves']:
            self.assertLess(engine.score(data, move['after']), engine.score(data, move['before']))

    def test_version_mismatch_rejected(self):
        data = fixture([match('one')]); result = engine.schedule(data)
        data['rest_minutes'] = 0
        self.assertFalse(engine.acceptance(data, result)['ready_for_export'])
        with self.assertRaises(ValueError): engine.improve(data, result)


class ExportAcceptance(unittest.TestCase):
    def test_offline_export_rechecks_instead_of_trusting_cached_pass(self):
        from export_event_tables import export_schedule
        from openpyxl import load_workbook
        data = fixture([match('one')])
        delayed = candidate(data, {'one': (3, 1)})
        delayed['acceptance'] = {'ready_for_export': True}
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'schedule.xlsx'
            with self.assertRaisesRegex(ValueError, 'local improvement'):
                export_schedule(data, delayed, path)
            self.assertFalse(path.exists())
            result = engine.improve(data, delayed)
            export_schedule(data, result, path)
            book = load_workbook(path)
            self.assertIn('优化验收', book.sheetnames)
            self.assertIn('约束来源', book.sheetnames)
            metrics = {row[0]: row[1:] for row in book['优化指标'].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(metrics['makespan'], ('60', '15'))

    def test_platform_file_adapter_uses_same_gate_without_network(self):
        from openpyxl import Workbook, load_workbook
        sys.path.insert(0, str(ROOT / 'ptty-skill/scripts'))
        spec = importlib.util.spec_from_file_location('synthetic_ptty_export', ROOT / 'ptty-skill/scripts/export_schedule.py')
        adapter = importlib.util.module_from_spec(spec); spec.loader.exec_module(adapter)
        m = match('one', platform_match_id='SYNTHETIC-PLATFORM-1', code='M1', project_name='Synthetic',
                  template_title='Synthetic', sides=[{'label': 'Synthetic A'}, {'label': 'Synthetic B'}])
        data = fixture([m]); delayed = candidate(data, {'one': (2, 1)})
        with tempfile.TemporaryDirectory() as temp:
            template, output, review = [Path(temp) / (n + '.xlsx') for n in ('template', 'output', 'review')]
            book = Workbook(); book.active.title = '赛事编排工作表'
            book.active.append(['日期', '时间', '场序', '第1号场地']); book.create_sheet('场次工作表'); book.save(template)
            with self.assertRaisesRegex(ValueError, 'local improvement'):
                adapter.export(data, delayed, template, output, review, ROOT / 'gamemaster-skill')
            self.assertFalse(output.exists())
            result = engine.improve(data, delayed)
            response = adapter.export(data, result, template, output, review, ROOT / 'gamemaster-skill')
            self.assertEqual(response['workbook_readback'], 'passed')
            self.assertEqual(load_workbook(output).sheetnames, ['赛事编排工作表', '场次工作表'])
            self.assertIn('早场阻塞', load_workbook(review).sheetnames)


if __name__ == '__main__':
    unittest.main(verbosity=2)
