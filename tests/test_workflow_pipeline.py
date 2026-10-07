"""Synthetic, offline cross-skill lifecycle regressions. Never connect to the platform."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'gamemaster-skill/scripts'))
sys.path.insert(0, str(ROOT / 'ptty-skill/scripts'))
import participant_roster
import prepare_event
import draw_engine
import bracket_engine
import schedule_engine
import results_engine
import export_competition_plan
import handoff
import validate_handoff
from openpyxl import load_workbook


def fixture(kind):
    count = 8 if kind == 'groups_knockout' else 4
    roster = {'reviewed': True, 'unresolved': [], 'exclusive': True, 'version': 1,
              'projects': [{'key': 'P', 'name': '合成男子单打', 'type': 'MS', 'expected_count': count}],
              'entries': [{'id': f'E{i}', 'status': 'active', 'project': 'P',
                           'source': '合成报名信息', 'seed': i if i <= 2 else 0,
                           'members': [{'id': f'A{i}', 'name': f'合成选手{i:02}',
                                        'unit': f'合成单位{(i - 1) % 2}', 'gender': 'M'}]}
                          for i in range(1, count + 1)]}
    p = {'id': 'P', 'sport': 'badminton', 'format': kind, 'entry_size': 1,
         'unit_policy': 'full_unit', 'duration_minutes': 15,
         'scoring': {'mode': 'points', 'best_of': 1, 'target': 21, 'win_by': 2, 'cap': 30},
         'ranking': {'criteria': [{'metric': 'wins', 'scope': 'all'}],
                     'restart_after_split': True, 'abnormal_results': 'adjudicate'}}
    if kind == 'knockout':
        p.update(seed_profile='badminton_bwf', third_place=True)
    else:
        p.update(group_sizes=[4, 4] if count == 8 else [4], seed_policy='snake')
        if kind == 'groups_knockout':
            p.update(third_place=True, advance_per_group=2,
                     knockout_slots=[{'group': g, 'rank': r} for g, r in [(1, 1), (2, 2), (2, 1), (1, 2)]])
    settings = {'confirmed': True, 'event_id': 'SYNTHETIC', 'event_name': '合成流程测试',
                'random_seed': 'synthetic-seed', 'projects': [p]}
    row = dict(id='P', group='合成公开组', type='MS', title='合成男子单打', stage=1, extra=1,
               entrants=count, groups=0 if kind == 'knockout' else count // 4,
               advance=2 if count == 8 else 4, start_rank=1,
               previous_groups=0, previous_start=0, previous_end=0, grab='', games=1,
               points=21, cap=30, scoring=1, format='TT' if kind == 'knockout' else 'XH',
               rotation=0, draw='0', group_sizes=p.get('group_sizes', []))
    rows = [row]
    if kind == 'groups_knockout':
        rows.append(dict(row, stage=2, entrants=4, groups=0, advance=4,
                         previous_groups=2, previous_start=1, previous_end=2, format='TT', draw='AB'))
    return roster, settings, {'confirmed': True, 'rows': rows}


def run_pipeline(kind, directory):
    directory = Path(directory)
    roster, settings, plan = fixture(kind)
    participant_roster.export(roster, directory / 'roster.xlsx')
    export_competition_plan.export(plan, directory / 'plan.xlsx')
    prepared = prepare_event.prepare(roster, settings)
    draw = draw_engine.generate(prepared['draw_config'])
    graph = bracket_engine.generate(prepare_event.attach_draw(prepared['draw_config'], prepared['bracket_base'], draw))
    bracket_engine.export_excel(graph, directory / 'bracket.xlsx')
    data = {'event_id': 'SYNTHETIC', 'courts': [1, 2, 3, 4], 'slot_minutes': 15,
            'sessions': [{'start': '2030-01-01T09:00', 'end': '2030-01-01T18:00', 'section': 1}],
            'rest_basis': 'elapsed', 'rest_minutes': 15, 'conflict_scope': 'possible',
            'matches': deepcopy(graph['matches']), 'outcome_disjoint_pairs': graph['outcome_disjoint_pairs'],
            'attempts': 4, 'random_seed': 7,
            'constraint_sources': {'*': {'kind': 'user', 'reference': 'Synthetic fixture specification'}}}
    schedule = schedule_engine.schedule(data)
    acceptance = schedule_engine.acceptance(data, schedule)
    assert acceptance['ready_for_export'], acceptance
    graph['synthetic'] = True
    graph['results'] = []
    # Winners form a strict total order so group ties do not need an invented tiebreak.
    while True:
        report = results_engine.evaluate(graph)
        if not report['pending_match_ids']:
            break
        ready = next((m for m in report['matches'] if m['status'] == 'pending' and all(m['entry_ids'])), None)
        assert ready, 'No legal next match; broken progression or unresolved ranking'
        winner = min(ready['entry_ids'])
        graph['results'].append({'match_id': ready['match_id'], 'status': 'completed', 'winner_id': winner,
                                 'score': {'games': [[21, 10] if ready['entry_ids'][0] == winner else [10, 21]]},
                                 'source': 'Synthetic result; not played'})
    results_engine.export_report(report, directory / 'results')
    for name, value in [('schedule.json', schedule), ('acceptance.json', acceptance), ('graph.json', graph)]:
        (directory / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    artifacts = [{'kind': p.suffix[1:], 'path': p.relative_to(directory).as_posix()}
                 for p in sorted(directory.rglob('*')) if p.is_file()]
    manifest = handoff.build({'event': {'local_id': 'SYNTHETIC', 'name': '合成流程测试', 'sport': 'badminton'},
                              'input_version': 'synthetic-v1', 'rules_version': 'synthetic-fixture-v1',
                              'artifacts': artifacts}, directory)
    target = directory / 'handoff.json'
    target.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    check = validate_handoff.validate(target)
    return graph, report, acceptance, check


class WorkflowPipeline(unittest.TestCase):
    def verify_pipeline(self, kind, expected_matches):
        with tempfile.TemporaryDirectory() as tmp:
            graph, report, acceptance, check = run_pipeline(kind, tmp)
            self.assertEqual(len(graph['matches']), expected_matches)
            self.assertEqual(len(report['matches']), expected_matches)
            self.assertEqual(report['unresolved_match_ids'], [])
            self.assertTrue(report['synthetic'])
            ranks = (report['groups'][0]['rows'] if kind == 'round_robin' else report['placements'])
            self.assertEqual({r['rank'] for r in ranks}, {1, 2, 3, 4})
            self.assertEqual(next(r['entry_id'] for r in ranks if r['rank'] == 1), 'E1')
            self.assertTrue(acceptance['hard_constraints']['complete'])
            self.assertFalse(acceptance['global_optimality_proven'])
            self.assertTrue(check['valid'])
            self.assertEqual(check['platform_event_id'], '')
            for path in Path(tmp).rglob('*.xlsx'):
                book = load_workbook(path, read_only=True)
                self.assertTrue(book.sheetnames)
                self.assertGreater(book.active.max_row, 1)
                book.close()
            # Exported artifacts cannot be edited silently after review.
            with (Path(tmp) / 'schedule.json').open('a') as file:
                file.write(' ')
            with self.assertRaisesRegex(ValueError, 'hash mismatch'):
                validate_handoff.validate(Path(tmp) / 'handoff.json')

    def test_round_robin_lifecycle(self):
        self.verify_pipeline('round_robin', 6)

    def test_knockout_lifecycle(self):
        self.verify_pipeline('knockout', 4)

    def test_groups_knockout_lifecycle(self):
        self.verify_pipeline('groups_knockout', 16)

    def test_team_aggregate_does_not_imply_team_roster_support(self):
        winner, statistics = results_engine.score_result({'rubbers': [2, 1]},
                                                        {'mode': 'team', 'win_target': 2, 'play_all': False})
        self.assertEqual(winner, 0)
        self.assertEqual(statistics['games'], [2, 1])
        roster, _, _ = fixture('knockout')
        roster['projects'][0]['type'] = 'MT'
        with self.assertRaisesRegex(ValueError, '专用模板'):
            participant_roster.validate(roster)


if __name__ == '__main__':
    unittest.main()
