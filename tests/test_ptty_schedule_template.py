"""Synthetic regressions for the platform's populated match-reference template."""
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'ptty-skill/scripts'))
from test_schedule_optimization import match, fixture, engine
spec = importlib.util.spec_from_file_location('ptty_schedule_reference_test', ROOT / 'ptty-skill/scripts/export_schedule.py')
adapter = importlib.util.module_from_spec(spec); spec.loader.exec_module(adapter)


def sample():
    m = match('M1', platform_match_id='SYNTHETIC-M1', platform_project_id='SYNTHETIC-P1',
              platform_format='TT', platform_display_code='SYNTHETIC-CODE1', round=1,
              code='LOCAL-M1', project_name='合成男单', template_title='合成标题',
              sides=[{'label': '合成甲'}, {'label': '合成乙'}])
    book = Workbook(); book.active.title = '赛事编排工作表'
    book.active.append(['日期', '时间', '场序', '第1号场地'])
    sheet = book.create_sheet('场次工作表')
    sheet.append(['项目ID', '项目全称', '赛事种类', '轮次'])
    sheet.append(['SYNTHETIC-P1', '合成男单', 'TT', '第1轮',
                  'SYNTHETIC-CODE1\r\n合成半决赛\r\n1-2\r\n合成A VS 合成B\r\n合成甲    合成乙\r\nSYNTHETIC-M1\r\n'])
    return book, m


class PopulatedTemplate(unittest.TestCase):
    def test_four_match_export_keeps_rest_gap_and_parallel_opposite_paths(self):
        matches = [match('semi-a', athletes=['a', 'b'], possible_athletes=['a', 'b'], round=1),
                   match('semi-b', athletes=['c', 'd'], possible_athletes=['c', 'd'], round=1)]
        for mid, kind in [('final', 'winner'), ('bronze', 'loser')]:
            matches.append(match(mid, round=2, possible_athletes=['a', 'b', 'c', 'd'],
                                 sides=[{'kind': kind, 'match_id': p, 'label': p + ' ' + kind}
                                        for p in ('semi-a', 'semi-b')]))
            matches[-1]['predecessors'] = ['semi-a', 'semi-b']
        book = Workbook(); book.active.title = '赛事编排工作表'
        book.active.append(['日期', '时间', '场序', '第1号场地', '第2号场地', '第3号场地', '第4号场地'])
        sheet = book.create_sheet('场次工作表')
        sheet.append(['项目ID', '项目全称', '赛事种类', '轮次'])
        for m in matches:
            m.update(platform_match_id='SYNTHETIC-' + m['id'], platform_project_id='SYNTHETIC-P',
                     platform_display_code='CODE-' + m['id'], platform_format='TT',
                     project_name='合成单打', code=m['id'], duration_minutes=30)
            m.setdefault('sides', [{'label': x} for x in m['athletes']])
            card = '\n'.join([m['platform_display_code'], '合成比赛', '?-?', '单位A VS 单位B',
                              '选手A 选手B', m['platform_match_id']])
            sheet.append(['SYNTHETIC-P', '合成单打', 'TT', '第' + str(m['round']) + '轮', card])
        data = fixture(matches, start='09:00', end='12:00', courts=(1, 2, 3, 4),
                       outcome_disjoint_pairs=[['final', 'bronze']])
        data.update(slot_minutes=30, rest_minutes=30, conflict_scope='possible')
        result = engine.schedule(data)
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory) / (n + '.xlsx') for n in ('template', 'output', 'review')]
            book.save(paths[0]); report = adapter.export(data, result, *paths, ROOT / 'gamemaster-skill')
            self.assertEqual(report['matches'], 4)
            grid = load_workbook(paths[1]).worksheets[0]
            starts = {}
            for row in grid.iter_rows(min_row=2, values_only=True):
                for cell in row[3:]:
                    if cell: starts[cell.strip().splitlines()[-1]] = row[1]
            self.assertEqual(starts, {'SYNTHETIC-semi-a': '09:00', 'SYNTHETIC-semi-b': '09:00',
                                      'SYNTHETIC-final': '10:00', 'SYNTHETIC-bronze': '10:00'})
            self.assertEqual(grid['B3'].value, '09:30')
            self.assertTrue(all(c.value is None for c in grid[3][3:]))
            self.assertEqual(result['metrics']['occupied_scenes'], 2)
            self.assertEqual(result['metrics']['scene_span'], 3)

    def test_doubles_card_and_repeated_header_full_export(self):
        book, m = sample(); sheet = book.worksheets[1]
        m['platform_event_type'] = 'MD'
        del m['template_title']  # Original platform card needs no fallback title.
        sheet['E2'] = sheet['E2'].value.replace('合成甲    合成乙', '合成甲    合成乙\r\n合成丙    合成丁')
        sheet.append(['项目ID', '项目全称', '赛事种类', '轮次'])
        data = fixture([m]); result = engine.schedule(data)
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory) / (n + '.xlsx') for n in ('template', 'output', 'review')]
            book.save(paths[0])
            report = adapter.export(data, result, *paths, ROOT / 'gamemaster-skill')
            self.assertEqual(report['verified_reference_cards'], 1)
            out = load_workbook(paths[1])
            self.assertIn('合成丙    合成丁', out.active['D2'].value)
            self.assertEqual(out.worksheets[1]['A3'].value, '项目ID')

    def test_seven_lines_require_doubles_and_exact_repeated_headers(self):
        book, m = sample(); sheet = book.worksheets[1]
        sheet['E2'] = sheet['E2'].value.replace('合成甲    合成乙', '合成甲    合成乙\n合成丙    合成丁')
        for kind in (None, 'MS', 'TEAM'):
            m['platform_event_type'] = kind
            with self.assertRaises(ValueError): adapter.match_cards(sheet, {'M1': m})
        m['platform_event_type'] = 'MD'
        sheet.append(['项目ID', '项目全称', '赛事种类', '轮次', 'unexpected'])
        with self.assertRaises(ValueError): adapter.match_cards(sheet, {'M1': m})

    def test_full_export_preserves_reference_and_uses_original_card(self):
        book, m = sample(); data = fixture([m]); result = engine.schedule(data)
        with tempfile.TemporaryDirectory() as directory:
            template, output, review = [Path(directory) / (name + '.xlsx') for name in ('template', 'output', 'review')]
            book.save(template)
            # Simulate an external producer encoding CRLF explicitly in XML.
            with ZipFile(template) as archive:
                contents = {n: archive.read(n) for n in archive.namelist()}
            xml = contents['xl/worksheets/sheet2.xml'].replace(b'&#13;', b'').replace(b'\r\n', b'\n')
            contents['xl/worksheets/sheet2.xml'] = xml.replace(b'\n', b'&#13;\n')
            with ZipFile(template, 'w') as archive:
                for name, value in contents.items(): archive.writestr(name, value)
            source = load_workbook(template)
            self.assertIn('\r\n', source.worksheets[1]['E2'].value)
            expected = [tuple(v.replace('\r\n', '\n') if isinstance(v, str) else v for v in row)
                        for row in source.worksheets[1].values]
            report = adapter.export(data, result, template, output, review, ROOT / 'gamemaster-skill')
            back = load_workbook(output)
            actual = [tuple(v.replace('\r\n', '\n') if isinstance(v, str) else v for v in row)
                      for row in back.worksheets[1].values]
            self.assertEqual(actual, expected)
            self.assertEqual(back.active['D2'].value.replace('\r\n', '\n'), expected[1][4])
            self.assertEqual(report['verified_reference_cards'], 1)

    def test_unbound_project_wrong_round_and_wrong_code_rejected(self):
        for cell, value in [('A2', 'SYNTHETIC-OTHER'), ('D2', '第2轮'), ('C2', 'XH')]:
            book, m = sample(); book.worksheets[1][cell] = value
            with self.assertRaises(ValueError): adapter.match_cards(book.worksheets[1], {'M1': m})
        book, m = sample(); del m['platform_project_id']
        with self.assertRaises(ValueError): adapter.match_cards(book.worksheets[1], {'M1': m})

    def test_unknown_duplicate_and_missing_match_rejected(self):
        for kind in ('unknown', 'duplicate', 'missing'):
            book, m = sample(); sheet = book.worksheets[1]
            if kind == 'unknown': sheet['E2'] = sheet['E2'].value.replace('SYNTHETIC-M1', 'SYNTHETIC-OTHER')
            elif kind == 'duplicate': sheet['F2'] = sheet['E2'].value
            else: sheet.delete_rows(2)
            with self.assertRaises(ValueError): adapter.match_cards(sheet, {'M1': m})

    def test_changed_schema_and_formula_rejected(self):
        for cell, value in [('A1', '另一个字段'), ('E2', '=1+1'), ('E2', 'incomplete card')]:
            book, m = sample(); book.worksheets[1][cell] = value
            with self.assertRaises(ValueError): adapter.match_cards(book.worksheets[1], {'M1': m})


if __name__ == '__main__':
    unittest.main()
