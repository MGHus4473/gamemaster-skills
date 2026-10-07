#!/usr/bin/env python3
"""Export a validated local schedule to the observed PTTY court-grid template."""
import argparse
import hashlib
import json
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core_bridge import load_core


def reference_values(sheet):
    # Java-produced templates may preserve CRLF; XML reserialization uses LF.
    return [tuple(v.replace('\r\n', '\n').replace('\r', '\n') if isinstance(v, str) else v for v in row)
            for row in sheet.values]


def match_cards(sheet, matches):
    """Read the observed populated match-reference sheet without rewriting it."""
    if not any(c.value is not None for row in sheet for c in row):
        return {}
    if sheet.merged_cells or any(c.data_type == 'f' for row in sheet for c in row):
        raise ValueError('Match-reference sheet must use unmerged static cells')
    header = [c.value for c in sheet[1]]
    if header[:4] != ['项目ID', '项目全称', '赛事种类', '轮次'] or any(v is not None for v in header[4:]):
        raise ValueError('Unrecognized match-reference headers')
    by_id = {m['platform_match_id']: m for m in matches.values()}
    cards = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(v in (None, '') for v in row):
            continue
        project, title, fmt, round_label = row[:4]
        if list(row[:4]) == header[:4] and all(v in (None, '') for v in row[4:]):
            continue  # The platform repeats the exact header between projects/stages.
        if not any(v not in (None, '') for v in row[4:]):
            raise ValueError('Match-reference row contains no matches')
        for value in row[4:]:
            if value in (None, ''):
                continue
            lines = value.rstrip('\r\n').splitlines() if isinstance(value, str) else []
            if len(lines) not in (6, 7) or lines[-1] not in by_id or lines[-1] in cards:
                raise ValueError('Unknown, duplicate or malformed match-reference card')
            m = by_id[lines[-1]]
            if len(lines) == 7 and m.get('platform_event_type') not in ('MD', 'WD', 'XD', 'SD'):
                raise ValueError('Two member rows require a verified doubles event type')
            expected = (m.get('platform_project_id'), m.get('project_name'), m.get('platform_format'),
                        '第' + str(m.get('round')) + '轮', m.get('platform_display_code'))
            if not all(expected) or (project, title, fmt, round_label, lines[0]) != expected:
                raise ValueError('Match-reference project, round or display code differs from verified mapping')
            cards[lines[-1]] = value
    if set(cards) != set(by_id):
        raise ValueError('Match-reference sheet must cover every scheduled match exactly once')
    return cards


def export(data, result, template, output, review, core_skill=None):
    output, review = Path(output), Path(review)
    if output.resolve() == review.resolve() or output.exists() or review.exists():
        raise ValueError('Choose two distinct, new output paths')
    if any(Path(p).suffix.lower() != '.xlsx' for p in (template, output, review)):
        raise ValueError('Scheduling template and both outputs must use .xlsx')
    engine = load_core('schedule_engine', core_skill)
    digest = hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    if result.get('event_id') != data['event_id'] or result.get('input_sha256') != digest:
        raise ValueError('Schedule result does not match event/input version')
    expected_slots = engine.prepare(data)[0]
    if result.get('slots') != expected_slots:
        raise ValueError('Schedule grid differs from confirmed sessions')
    report = engine.validate(data, result)
    if not report['complete']:
        raise ValueError('Refusing formal export: ' + str(report))
    audit = engine.acceptance(data, result)
    if not audit['ready_for_export']:
        raise ValueError('Schedule needs source confirmation or local improvement: ' +
                         json.dumps({k: audit[k] for k in ('sources_complete', 'local_search')}, ensure_ascii=False))
    matches = {m['id']: m for m in data['matches']}
    if any(not m.get('platform_match_id') for m in matches.values()):
        raise ValueError('Real platform IDs required for PTTY import')
    ids = [m['platform_match_id'] for m in matches.values()]
    if len(set(ids)) != len(ids):
        raise ValueError('Duplicate platform match IDs')
    wb = load_workbook(template)
    if wb.sheetnames != ['赛事编排工作表', '场次工作表']:
        raise ValueError('Unrecognized scheduling template')
    cards = match_cards(wb.worksheets[1], matches)
    secondary_values = reference_values(wb.worksheets[1])
    ws = wb.worksheets[0]
    if [ws.cell(1, n).value for n in (1, 2, 3)] != ['日期', '时间', '场序']:
        raise ValueError('Unrecognized fixed columns')
    if ws.merged_cells:
        raise ValueError('Merged template needs a separate adapter')
    styles = [copy(ws.cell(2, n)._style) for n in range(1, 5)]
    headers = [copy(ws.cell(1, n)._style) for n in range(1, 5)]
    width = ws.column_dimensions['D'].width
    ws.delete_rows(2, ws.max_row)
    ws.delete_cols(4, max(0, ws.max_column - 3))
    for col, court in enumerate(data['courts'], 4):
        ws.cell(1, col, f'第{court}号场地')._style = copy(headers[3])
        ws.column_dimensions[get_column_letter(col)].width = width
    cells = {(r['slot_index'], r['court']): r for r in result['assignments']}
    for slot in result['slots']:
        row = [slot['start'][:10], slot['start'][11:16], str(slot['scene'])]
        for court in data['courts']:
            rec = cells.get((slot['index'], court))
            if rec is None:
                row.append('')
                continue
            m = matches[rec['match_id']]
            if m['platform_match_id'] in cards:
                row.append(cards[m['platform_match_id']])
            else:
                a, b = m['sides']
                text = [m['code'], m['template_title'],
                        a.get('position_label', '?') + '-' + b.get('position_label', '?'),
                        a.get('club', '') + '   VS  ' + b.get('club', ''),
                        a['label'] + '    ' + b['label'], m['platform_match_id'], '']
                row.append('\r\n'.join(text))
        for col, value in enumerate(row, 1):
            c = ws.cell(slot['index'] + 2, col, value)
            c._style = copy(styles[min(col - 1, 3)])
            c.data_type = 's'
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[slot['index'] + 2].height = 110
    ws.freeze_panes = 'D2'
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    ws.print_area = f'A1:{get_column_letter(3 + len(data["courts"]))}{len(result["slots"]) + 1}'
    output.parent.mkdir(parents=True, exist_ok=True)
    review.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    # Re-open the exact upload artifact; compare time, court, scene and IDs.
    check = load_workbook(output, data_only=False)
    if reference_values(check.worksheets[1]) != secondary_values:
        raise ValueError('Saved match-reference sheet changed')
    actual = {}
    for row in check.worksheets[0].iter_rows(min_row=2):
        for i, c in enumerate(row[3:]):
            if c.value:
                mid = str(c.value).strip().splitlines()[-1]
                if mid in actual:
                    raise ValueError('Repeated ID in output')
                actual[mid] = (row[0].value + 'T' + row[1].value, int(row[2].value), data['courts'][i])
    expected = {matches[r['match_id']]['platform_match_id']: (r['start'], r['scene'], r['court'])
                for r in result['assignments']}
    if actual != expected:
        raise ValueError('Saved workbook differs from validated schedule')
    book = Workbook()
    sheet = book.active
    sheet.title = '编排明细'
    sheet.append(['日期', '开始', '预计结束', '场序', '小节', '场地', '项目', '轮次', '名次范围',
                  '赛号', '对阵一', '对阵二', '前置场次', '系统场次ID'])
    for r in result['assignments']:
        m = matches[r['match_id']]
        sheet.append([r['start'][:10], r['start'][11:], r['end'][11:], r['scene'], r['section'],
                      r['court'], m['project_name'], m.get('round'), m.get('rank'), m['code'],
                      m['sides'][0]['label'], m['sides'][1]['label'],
                      '、'.join(matches[p]['code'] for p in m.get('predecessors', [])), m['platform_match_id']])
    config = book.create_sheet('参数与核验')
    config.append(['项目', '内容'])
    for key in ('event_id', 'slot_minutes', 'courts', 'sessions', 'objective', 'priority_projects',
                'conflict_scope', 'rest_basis', 'rest_minutes', 'rest_scenes', 'turnover_minutes'):
        config.append([key, json.dumps(data.get(key), ensure_ascii=False)])
    config.append(['核验', json.dumps(report, ensure_ascii=False)])
    config.append(['待定选手', '已确认名单模式：赛果确定后重新检查跨项目冲突；所有前后轮保留休息。'
                  if data['conflict_scope'] == 'known' else '按所有可能选手保守检查；互斥分支单独核对。'])
    config.append(['预计时长', '所有时间为编排预计；实际超时后按实际结束时间复核休息并调整未开赛场次。'])
    config.append(['测试种子', data.get('seed_note', '')])
    for title, rows in engine.audit_tables(data, result, audit).items():
        audit_sheet = book.create_sheet(title)
        for row in rows:
            audit_sheet.append(row)
    for s in book:
        s.freeze_panes = 'A2'
        s.auto_filter.ref = s.dimensions
        for cell in s[1]:
            cell.fill = PatternFill('solid', fgColor='24476A')
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        for row in s.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            s.row_dimensions[row[0].row].height = 36
        for col in range(1, s.max_column + 1):
            s.column_dimensions[get_column_letter(col)].width = 16 if col < 10 else 30
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.print_title_rows = '1:1'
    config.column_dimensions['A'].width = 24
    config.column_dimensions['B'].width = 100
    book.save(review)
    return {'matches': len(actual), 'rows': len(result['slots']), 'courts': len(data['courts']),
            'import': str(output), 'review': str(review), 'workbook_readback': 'passed',
            'verified_reference_cards': len(cards), 'preserved_reference_sheet': True}


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    for name in ('input', 'result', 'template', 'output', 'review'):
        p.add_argument(name, type=Path)
    p.add_argument('--core-skill', type=Path)
    a = p.parse_args()
    print(json.dumps(export(json.loads(a.input.read_text()), json.loads(a.result.read_text()),
                            a.template, a.output, a.review, a.core_skill), ensure_ascii=False))
