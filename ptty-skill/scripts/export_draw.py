#!/usr/bin/env python3
"""Export a validated offline draw to an existing PTTY .xls draw template."""
import argparse
import json
from pathlib import Path

import xlrd
from xlutils.copy import copy as copy_xls
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core_bridge import load_core


def require(ok, message):
    if not ok:
        raise ValueError(message)

HEADERS = ['项目ID','阶段','附加','组别','项目名称','项目类型','赛事种类','组号/轮次号','位置号','2分区','4分区','8分区','16分区','队伍名称','姓名','性别','队内技术号','种子号','前阶段名次','识别码']


def export(config, result, mapping, template, output, core_skill=None):
    engine = load_core('draw_engine', core_skill)
    output = Path(output)
    require(Path(template).suffix.lower() == '.xls' and output.suffix.lower() == '.xls',
            '当前适配器输入模板和输出必须为真实.xls文件')
    require(not output.exists() and not output.with_name(output.stem+'_审阅.xlsx').exists(), '输出已存在，请换新版本')
    require(result['event_id'] == config['event_id'] and result['input_sha256'] == engine.digest(config), '抽签结果与输入版本不一致')
    projects = {p['id']: p for p in config['projects']}
    require(set(projects) == {p['project_id'] for p in result['projects']}, '输出项目不一致')
    for p in result['projects']:
        engine.validate_project(projects[p['project_id']], p)
    ids = {(x['project_id'], x['entry_id']): x for x in mapping}
    require(len(ids) == len(mapping) and len({x['XMNM'] for x in mapping}) == len(mapping), '报名项映射重复')
    require(set(ids) == {(p['id'], e['id']) for p in config['projects'] for e in p['entries']},
            '报名项映射有额外项目或遗漏')
    project_map = {}
    for item in mapping:
        local_id = item['project_id']
        platform_id = item.get('platform_project_id', local_id)
        require(isinstance(platform_id, str) and bool(platform_id), '缺少目标平台项目ID')
        require(local_id not in project_map or project_map[local_id] == platform_id,
                '同一本地项目映射到多个平台项目')
        project_map[local_id] = platform_id
    require(len(set(project_map.values())) == len(project_map), '多个本地项目不能共用目标平台项目ID')
    local_for_platform = {value: key for key, value in project_map.items()}
    source = xlrd.open_workbook(str(template), formatting_info=True)
    require(source.sheet_names() == ['抽签工作表','人员排名表（第1阶段）'], '当前适配器仅支持已验证的第一阶段20列.xls模板')
    s = source.sheet_by_index(0); rank = source.sheet_by_index(1)
    require(s.ncols == 20 and s.row_values(1) == HEADERS, '导签模板字段发生变化')
    entrants = {str(rank.cell_value(i,9)): (i,rank.row_values(i)) for i in range(1,rank.nrows)}
    require(set(entrants) == {m['XMNM'] for m in mapping}, '模板报名项与映射不一致')
    require(len(entrants) == rank.nrows-1, '模板报名项ID重复')
    target = copy_xls(source)
    sheet = target.get_sheet(0); ranksheet=target.get_sheet(1)
    template_slots = {}; bye_slots = {}
    for i in range(2,s.nrows):
        platform_id = str(s.cell_value(i,0))
        require(platform_id in local_for_platform, '模板项目未绑定到本地项目：' + platform_id)
        xmid = local_for_platform[platform_id]; fmt=projects[xmid]['format']
        require(str(s.cell_value(i,1))=='1' and str(s.cell_value(i,2))=='1', '当前适配器只验证阶段1附加1')
        pos=int(s.cell_value(i,8))
        if fmt=='knockout':group=1
        else:
            raw=str(s.cell_value(i,7)).removesuffix('组')
            require(raw.isdecimal(), '循环组号须为数字；字母组号需另行映射')
            group=int(raw)
        key=(xmid,group,pos)
        require(key not in template_slots, '模板签位重复')
        template_slots[key]=i
        if s.cell_value(i,14)=='轮空':bye_slots.setdefault(xmid,[]).append(pos)
    for p in result['projects']:
        if p['format']=='knockout':
            require(sorted(bye_slots.get(p['project_id'],[])) == p['byes'], '离线轮空与已有平台赛程不一致，禁止移动轮空')
    written=set(); summary=[]
    def put(ws,row,col,value):
        cell=ws.row(row)._Row__cells[col]
        style=cell.xf_idx
        ws.write(row,col,value)
        ws.row(row)._Row__cells[col].xf_idx=style
    for project in result['projects']:
        for a in project['assignments']:
            m=ids[(project['project_id'],a['entry_id'])]
            require(m['seed']==a['seed'] and m['club']==a['club'] and m['name']==a['name'], '人员映射与抽签记录不符')
            source_row,person=entrants[m['XMNM']]
            values=person[3:10]
            require(values[0] == a['club'] and values[1] == a['name'], '模板姓名/队名与报名项不一致')
            values[4]=str(a['seed']) if a['seed'] else ''
            key=(project['project_id'],a['group'],a['position'])
            require(key in template_slots and key not in written, '导出位置不存在或重复')
            row=template_slots[key]
            require(s.cell_value(row,14)!='轮空', '不能向轮空位置填人')
            for col,value in enumerate(values,13):put(sheet,row,col,value)
            put(ranksheet,source_row,7,values[4])
            written.add(key)
            summary.append([project['name'],a['group'],a['position'],a['seed'] or '',a['name'],a['club'],a['entry_id'],m['XMNM'],'临时测试种子（不代表实力排名）' if a['seed'] and result.get('test_seeds') else ''])
    require(len(written)==sum(len(p['entries']) for p in projects.values()), '导出条数不足')
    output.parent.mkdir(parents=True,exist_ok=True);target.save(str(output))
    back=xlrd.open_workbook(str(output));b=back.sheet_by_index(0)
    require(back.sheet_names()==source.sheet_names() and b.nrows==s.nrows, '模板工作表或行数改变')
    for i in range(s.nrows):
        cols=range(20) if i<2 or s.cell_value(i,14)=='轮空' else range(13)
        require(all(b.cell_value(i,c)==s.cell_value(i,c) for c in cols), '固定模板结构或轮空被更改')
    observed={str(b.cell_value(i,19)) for i in range(2,b.nrows) if b.cell_value(i,14)!='轮空'}
    require(observed==set(entrants), '导出后识别码有漏项')
    review=Workbook();r=review.active;r.title='抽签结果'
    r.append(['项目','组号','签位','种子号','姓名或组合','完整队名','本地报名项ID','跑兔识别码','备注'])
    for row in summary:r.append(row)
    r.freeze_panes='A2';r.auto_filter.ref=r.dimensions
    widths=[12,8,8,10,27,50,18,26,34]
    for i,w in enumerate(widths,1):r.column_dimensions[chr(64+i)].width=w
    for c in r[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='215E75')
    for row in r.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Arial',size=11);c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(c.value,str):c.data_type='s'
    r2=review.create_sheet('抽签记录')
    for row in [['事项','记录'],['随机种子',result['random_seed']],['输入SHA256',result['input_sha256']],['测试种子',str(result.get('test_seeds',False))],['规则','种子按层落位；完整队名在每层分区人数差不超过1；不合并一队二队'],['验证',f'{len(written)}个报名项；各分区数量见抽签结果JSON']]:r2.append(row)
    for row in r2:
        for cell in row:
            if isinstance(cell.value, str):cell.data_type='s'
    r2.column_dimensions['A'].width=18;r2.column_dimensions['B'].width=95
    review.save(output.with_name(output.stem+'_审阅.xlsx'))
    return {'passed':True,'entries':len(written),'preserved_template_fixed_columns':True,'preserved_byes':True,'mapping_unique':True,'project_id_mapping':project_map}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    for name in ['input','result','mapping','template','output']:p.add_argument(name,type=Path)
    p.add_argument('--core-skill', type=Path)
    a=p.parse_args();load=lambda path:json.loads(path.read_text(encoding='utf-8'))
    print(json.dumps(export(load(a.input),load(a.result),load(a.mapping),a.template,a.output,a.core_skill),ensure_ascii=False))

if __name__=='__main__':main()
