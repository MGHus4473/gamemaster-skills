#!/usr/bin/env python3
"""Offline, explicit-rule results validation, standings and portable exports.

No network or platform IDs. JSON/CSV use the standard library; XLSX uses openpyxl.
Input schema and limitations: ../references/results.md.
"""
import argparse
from copy import deepcopy
import csv
from fractions import Fraction
from hashlib import sha256
import json
from pathlib import Path


FINAL = {'completed', 'walkover', 'retired'}
STATUS = {'pending': '未赛', 'completed': '完赛', 'walkover': '弃权',
          'retired': '退赛', 'cancelled': '取消'}
METRICS = {'wins', 'match_points', 'game_diff', 'game_ratio', 'point_diff', 'point_ratio'}


def require(ok, message):
    if not ok:
        raise ValueError(message)


def integer(value, label, minimum=0):
    require(type(value) is int and value >= minimum, f'{label}: integer >= {minimum} required')
    return value


def pair(value, label='score'):
    require(isinstance(value, list) and len(value) == 2, f'{label}: two scores required')
    return [integer(x, label) for x in value]


def terminal(points, target, win_by, cap=None):
    a, b = points
    high, low = max(a, b), min(a, b)
    return high >= target and (high - low >= win_by or (cap is not None and high == cap and high > low))


def point_winner(points, target, win_by=2, cap=None, initial_points=None):
    """Require an exact final score: reject play continuing after a game was won."""
    points = pair(points)
    initial = pair(initial_points if initial_points is not None else [0, 0], 'initial_points')
    require(all(points[i] >= initial[i] for i in range(2)), 'Score below confirmed initial points')
    require(cap is None or max(points) <= cap, 'Score exceeds cap')
    require(terminal(points, target, win_by, cap), 'Incomplete or tied game score')
    winner = 0 if points[0] > points[1] else 1
    previous = points[:]
    previous[winner] -= 1
    require(not terminal(previous, target, win_by, cap), 'Score continued after game was already won')
    return winner


def validate_rule(rule):
    require(isinstance(rule, dict), 'Explicit scoring rule required')
    mode = rule.get('mode')
    require(mode in ('points', 'tennis', 'team'), 'Unknown scoring mode')
    require(mode == 'points' or 'initial_points' not in rule, 'initial_points supported only by points mode')
    if mode == 'team':
        integer(rule.get('win_target'), 'team win_target', 1)
        require(type(rule.get('play_all')) is bool, 'team play_all must be explicit')
        if rule['play_all']:
            count = integer(rule.get('rubber_count'), 'team rubber_count', 1)
            require(count % 2 == 1 and rule['win_target'] == count // 2 + 1,
                    'All-rubbers mode supports odd count with majority winner')
        return
    best = integer(rule.get('best_of'), 'best_of', 1)
    require(best % 2 == 1, 'best_of must be odd')
    if mode == 'points':
        target = integer(rule.get('target'), 'target', 1)
        integer(rule.get('win_by'), 'win_by', 1)
        cap = rule.get('cap')
        if cap is not None:
            require(integer(cap, 'cap', 1) >= target, 'cap below target')
        initial = pair(rule.get('initial_points', [0, 0]), 'initial_points')
        require(cap is None or max(initial) < cap, 'Initial score reaches/exceeds game cap')
        require(not terminal(initial, target, rule['win_by'], cap), 'Initial score already wins the game')
    else:
        integer(rule.get('games_to_win'), 'games_to_win', 1)
        integer(rule.get('set_win_by'), 'set_win_by', 1)
        for key in ('tiebreak_at', 'final_tiebreak_at'):
            if rule.get(key) is not None:
                integer(rule[key], key, 1)
        integer(rule.get('tiebreak_target'), 'tiebreak_target', 1)
        if rule.get('final_tiebreak_target') is not None:
            integer(rule['final_tiebreak_target'], 'final_tiebreak_target', 1)
        if rule.get('match_tiebreak_target') is not None:
            integer(rule['match_tiebreak_target'], 'match_tiebreak_target', 1)


def tennis_set(item, rule, deciding=False):
    require(isinstance(item, dict) and 'games' in item, 'Tennis set needs games')
    require(set(item) <= {'games', 'tiebreak'}, 'Unknown or conflicting tennis set fields')
    games = pair(item['games'], 'set games')
    at = rule.get('final_tiebreak_at', rule.get('tiebreak_at')) if deciding else rule.get('tiebreak_at')
    target = rule.get('final_tiebreak_target', rule['tiebreak_target']) if deciding else rule['tiebreak_target']
    if at is not None and sorted(games) == [at, at + 1]:
        winner = point_winner(item.get('tiebreak'), target, 2)
        require(games[winner] == at + 1, 'Tiebreak winner disagrees with set games')
        return winner
    require('tiebreak' not in item, 'Tiebreak score attached to non-tiebreak set')
    if at is not None:
        require(min(games) < at, 'Set passed configured tiebreak trigger')
    return point_winner(games, rule['games_to_win'], rule['set_win_by'])


def score_result(score, rule, complete=True):
    """Return winning side (None for incomplete prefix) and ranking quantities.

    Tennis game_* statistics mean sets, point_* mean ordinary games. A match
    tiebreak counts as one set, zero ordinary games; ranking must accept that
    policy explicitly. The full original structure is always preserved.
    """
    validate_rule(rule)
    require(isinstance(score, dict), 'Structured score required')
    allowed = {'points': {'games', 'partial_game'}, 'tennis': {'sets', 'partial_set'}, 'team': {'rubbers'}}
    require(set(score) <= allowed[rule['mode']], 'Unknown or mixed scoring fields')
    if rule['mode'] == 'team':
        rubbers = pair(score.get('rubbers'), 'team rubbers')
        require(complete, 'Partial team score needs separate adjudication')
        target = rule['win_target']
        require(max(rubbers) >= target and rubbers[0] != rubbers[1], 'Team tie not decided')
        if rule['play_all']:
            require(sum(rubbers) == rule['rubber_count'], 'Incomplete all-rubbers team tie')
        else:
            require(max(rubbers) == target and min(rubbers) < target, 'Team tie continued after victory')
        return (0 if rubbers[0] > rubbers[1] else 1), {'games': rubbers, 'points': [0, 0]}
    wins, points = [0, 0], [0, 0]
    needed = rule['best_of'] // 2 + 1
    key = 'games' if rule['mode'] == 'points' else 'sets'
    units = score.get(key, [])
    require(isinstance(units, list), f'{key} must be a list')
    for item in units:
        require(max(wins) < needed, 'Extra game/set after match victory')
        if rule['mode'] == 'points':
            winner = point_winner(item, rule['target'], rule['win_by'], rule.get('cap'), rule.get('initial_points'))
            quantities = item
        elif 'match_tiebreak' in item:
            require(rule.get('match_tiebreak_target') and wins == [needed - 1, needed - 1],
                    'Match tiebreak only replaces a configured deciding set')
            require(set(item) == {'match_tiebreak'}, 'Match tiebreak cannot also contain set games')
            winner = point_winner(item['match_tiebreak'], rule['match_tiebreak_target'], 2)
            quantities = [0, 0]
        else:
            deciding = wins == [needed - 1, needed - 1]
            require(not (deciding and rule.get('match_tiebreak_target')), 'Deciding set must be match tiebreak')
            winner = tennis_set(item, rule, deciding)
            quantities = item['games']
        wins[winner] += 1
        points = [points[i] + quantities[i] for i in range(2)]
    winner = wins.index(needed) if needed in wins else None
    require(not complete or winner is not None, 'Match has no winning majority')
    require(complete or winner is None, 'Retirement prefix already contains a completed match')
    partial = score.get('partial_game')
    if partial is not None:
        require(not complete and rule['mode'] == 'points', 'partial_game only allowed in retired point match')
        partial = pair(partial)
        initial = rule.get('initial_points', [0, 0])
        require(all(partial[i] >= initial[i] for i in range(2)), 'Partial score below confirmed initial points')
        require(rule.get('cap') is None or max(partial) < rule['cap'], 'Partial score reaches/exceeds game cap')
        require(not terminal(partial, rule['target'], rule['win_by'], rule.get('cap')), 'Partial game already terminal')
        points = [points[i] + partial[i] for i in range(2)]
    partial_set = score.get('partial_set')
    if partial_set is not None:
        require(not complete and rule['mode'] == 'tennis', 'partial_set only allowed in retired tennis')
        # Preserve unfinished set; do not pretend to validate point-by-point history.
        require(isinstance(partial_set, dict), 'partial_set must be structured')
        if 'match_tiebreak' in partial_set:
            require(rule.get('match_tiebreak_target') and wins == [needed - 1, needed - 1], 'Invalid partial match tiebreak')
            partial = pair(partial_set['match_tiebreak'])
            require(not terminal(partial, rule['match_tiebreak_target'], 2), 'Partial match tiebreak already terminal')
        else:
            partial = pair(partial_set.get('games'), 'partial set games')
            deciding = wins == [needed - 1, needed - 1]
            require(not (deciding and rule.get('match_tiebreak_target')), 'Expected partial match tiebreak')
            at = rule.get('final_tiebreak_at', rule.get('tiebreak_at')) if deciding else rule.get('tiebreak_at')
            require(not terminal(partial, rule['games_to_win'], rule['set_win_by']), 'Partial set already terminal')
            require(at is None or max(partial) <= at, 'Partial set passed tiebreak trigger')
            if 'tiebreak' in partial_set:
                require(at is not None and partial == [at, at], 'Partial tiebreak at wrong games score')
                target = rule.get('final_tiebreak_target', rule['tiebreak_target']) if deciding else rule['tiebreak_target']
                require(not terminal(pair(partial_set['tiebreak']), target, 2), 'Partial tiebreak already terminal')
            points = [points[i] + partial[i] for i in range(2)]
    return winner, {'games': wins, 'points': points}


def index(rows, key='id'):
    require(isinstance(rows, list), 'Rows must be a list')
    require(all(isinstance(x.get(key), str) and x[key] for x in rows), f'Non-empty {key} required')
    require(len({x[key] for x in rows}) == len(rows), f'Duplicate {key}')
    return {x[key]: x for x in rows}


def rank_group(group, matches):
    """Rank a complete, explicit round robin, preserving unresolved tie blocks."""
    ids = group['entry_ids']
    require(len(ids) == len(set(ids)), 'Duplicate group entry')
    excluded = group.get('excluded_entries', {})
    require(set(excluded) <= set(ids) and all(excluded.values()), 'Excluded entries need reasons')
    active = [i for i in ids if i not in excluded]
    require(len(active) >= 2, 'At least two eligible entries required to calculate group standings')
    policy = group.get('ranking', {})
    criteria = policy.get('criteria')
    require(isinstance(criteria, list) and criteria, 'Explicit ranking criteria required')
    require(type(policy.get('restart_after_split')) is bool, 'ranking.restart_after_split must be explicit')
    for c in criteria:
        require(c.get('metric') in METRICS and c.get('scope') in ('all', 'tied'), 'Invalid ranking criterion')
    usable, counts, pending = [], {}, []
    for m in matches:
        sides = m['entry_ids']
        require(len(sides) == 2 and all(i in ids for i in sides), 'Group match has unresolved or foreign entries')
        if any(i in excluded for i in sides):
            continue  # Explicit whole-group annulment, not a silent retirement assumption.
        k = tuple(sorted(sides))
        counts[k] = counts.get(k, 0) + 1
        if m['status'] not in FINAL:
            pending.append(m['match_id'])
        else:
            usable.append(m)
    meetings = integer(group.get('expected_meetings', 1), 'expected_meetings', 1)
    require(all(counts.get(tuple(sorted((a, b))), 0) == meetings for n, a in enumerate(active) for b in active[n + 1:]),
            'Round robin pair coverage differs from expected_meetings')
    abnormal = [m['match_id'] for m in usable if m['status'] != 'completed']
    abnormal_policy = policy.get('abnormal_results', 'adjudicate')
    require(abnormal_policy in ('adjudicate', 'wins_only', 'recorded_score', 'awarded_score'), 'Unknown abnormal results policy')
    issue = []
    if pending:
        issue.append('组内比赛未全部完成')
    if abnormal and abnormal_policy == 'adjudicate':
        issue.append('弃权/退赛的循环排名处理待裁定')
    if abnormal and abnormal_policy == 'wins_only':
        require(all(c['metric'] in ('wins', 'match_points') for c in criteria), 'wins_only cannot use score-derived tie breakers')
    if abnormal and abnormal_policy == 'recorded_score':
        require(policy.get('abnormal_rule_reference'), 'recorded_score requires confirmed regulation reference')
        if any(c['metric'] not in ('wins', 'match_points') for c in criteria):
            require(all(m.get('statistics') for m in usable), 'recorded_score cannot treat missing actual scores as zero')
    if abnormal and abnormal_policy == 'awarded_score':
        require(policy.get('abnormal_rule_reference'), 'awarded_score requires confirmed ruling reference')
        require(all(m['status'] == 'completed' or (m['status'] in ('walkover', 'retired') and m.get('awarded_statistics')) for m in usable),
                'awarded_score requires a separately documented administrative score for every abnormal result')
    if any(m.get('scoring_mode') == 'tennis' for m in usable) and any(c['metric'].startswith(('game_', 'point_')) for c in criteria):
        require(policy.get('tennis_stat_policy') == 'sets_and_ordinary_games_excluding_match_tiebreak',
                'Explicit tennis ranking score treatment required')
    def stats(pool):
        s = {i: {'played': 0, 'wins': 0, 'losses': 0, 'match_points': 0,
                 'games_for': 0, 'games_against': 0, 'points_for': 0, 'points_against': 0} for i in pool}
        for m in usable:
            if not set(m['entry_ids']) <= set(pool):
                continue
            for side, i in enumerate(m['entry_ids']):
                win = i == m['winner_id']
                s[i]['played'] += 1
                s[i]['wins' if win else 'losses'] += 1
                if any(c['metric'] == 'match_points' for c in criteria):
                    scoring = policy.get('match_points', {})
                    require('win' in scoring and 'loss' in scoring, 'Win/loss match points must be explicit')
                    award = scoring['win'] if win else scoring.get(m['status'] + '_loss', scoring['loss'])
                    s[i]['match_points'] += integer(award, 'match points')
                statistics = m.get('awarded_statistics') if abnormal_policy == 'awarded_score' and m['status'] != 'completed' else m.get('statistics')
                if statistics:
                    for dst, src in (('games', 'games'), ('points', 'points')):
                        s[i][dst + '_for'] += statistics[src][side]
                        s[i][dst + '_against'] += statistics[src][1 - side]
        return s
    overall = stats(active)
    trace = []
    def value(s, metric):
        if metric in ('wins', 'match_points'):
            return s[metric]
        prefix, operation = metric.split('_')
        a, b = s[prefix + 's_for'], s[prefix + 's_against']
        if operation == 'diff':
            return a - b
        # Exact ratio; zero denominator wins over every finite ratio, 0/0 stays 0.
        return (1, Fraction(0)) if b == 0 and a else (0, Fraction(a, b) if b else Fraction(0))
    def split(block, step):
        if len(block) == 1 or step == len(criteria):
            return [block]
        c = criteria[step]
        st = overall if c['scope'] == 'all' else stats(block)
        buckets = {}
        for i in block:
            buckets.setdefault(value(st[i], c['metric']), []).append(i)
        ordered = [buckets[v] for v in sorted(buckets, reverse=True)]
        trace.append({'entry_ids': block, 'criterion': c, 'blocks': ordered})
        if len(ordered) == 1:
            return split(block, step + 1)
        nxt = 0 if policy['restart_after_split'] else step + 1
        return [b for sub in ordered for b in split(sub, nxt)]
    blocks = split(active, 0) if not issue else [active]
    used_decisions = set()
    decisions = group.get('tie_decisions', [])
    if not issue:
        resolved_blocks = []
        for block in blocks:
            applicable = [(n, d) for n, d in enumerate(decisions)
                          if len(block) > 1 and set(d.get('order', [])) == set(block)]
            require(len(applicable) <= 1, 'Multiple decisions for the same unresolved tie')
            if applicable:
                n, decision = applicable[0]
                require(decision.get('confirmed') is True and decision.get('reason'),
                        'Tie decision requires confirmed order and reason')
                require(len(decision['order']) == len(block), 'Tie decision contains duplicate entries')
                used_decisions.add(n)
                trace.append({'entry_ids': block, 'decision': deepcopy(decision)})
                resolved_blocks.extend([[i] for i in decision['order']])
            else:
                resolved_blocks.append(block)
        blocks = resolved_blocks
    require(len(used_decisions) == len(decisions), 'Tie decision does not match a currently unresolved complete tie')
    rows, position = [], 1
    for block in blocks:
        for i in block:
            rows.append({'entry_id': i, 'rank': position if not issue else None,
                         'tied': len(block) > 1, 'status': '待裁定' if issue or len(block) > 1 else '已核定',
                         **overall[i]})
        position += len(block)
    tied = any(len(b) > 1 for b in blocks)
    return {'id': group['id'], 'project_id': group['project_id'], 'rows': rows, 'trace': trace,
            'complete': not issue, 'resolved': not issue and not tied, 'issues': issue,
            'excluded_entries': excluded, 'pending_match_ids': pending,
            'status': '待裁定' if issue or tied else '已核定'}


def evaluate(data):
    """Validate all results against a match DAG, calculate groups and placements."""
    projects = index(data.get('projects', []))
    entries = index(data.get('entries', []))
    matches = index(data.get('matches', []))
    groups = index(data.get('groups', []))
    results = index(data.get('results', []), 'match_id')
    require(set(results) <= set(matches), 'Result references an unknown match')
    for p in projects.values():
        require(p.get('sport') in ('badminton', 'table_tennis', 'tennis', 'pickleball'), 'Explicit supported sport required')
        validate_rule(p.get('scoring'))
        for rule in p.get('scoring_by_stage', {}).values():
            validate_rule(rule)
    for e in entries.values():
        require(e.get('project_id') in projects, 'Entry project unknown')
    for g in groups.values():
        require(g.get('project_id') in projects, 'Group project unknown')
        require(set(g.get('entry_ids', [])) <= set(entries), 'Group entry unknown')
        require(all(entries[i]['project_id'] == g['project_id'] for i in g['entry_ids']), 'Group mixes projects')
        require(set(g.get('match_ids', [])) <= set(matches), 'Group match unknown')
        require(len(g['match_ids']) == len(set(g['match_ids'])), 'Duplicate group match')
    cache, gcache, visiting = {}, {}, set()
    def group_result(gid):
        require(gid in groups, 'Unknown group source')
        if gid in gcache:
            return gcache[gid]
        token = ('group', gid)
        require(token not in visiting, 'Cyclic group/match dependency')
        visiting.add(token)
        g = groups[gid]
        gm = [match_result(i) for i in g['match_ids']]
        require(all(m['project_id'] == g['project_id'] for m in gm), 'Group match project mismatch')
        gcache[gid] = rank_group(g, gm)
        visiting.remove(token)
        return gcache[gid]
    def resolve(side, pid):
        kind = side.get('kind')
        if kind == 'entry':
            eid = side.get('entry_id')
            require(eid in entries and entries[eid]['project_id'] == pid, 'Unknown/foreign entry source')
            return eid
        if kind in ('winner', 'loser'):
            dep = match_result(side.get('match_id'))
            require(dep['project_id'] == pid, 'Cross-project match dependency')
            return dep.get('winner_id' if kind == 'winner' else 'loser_id')
        if kind == 'group_rank':
            g = group_result(side.get('group_id'))
            require(g['project_id'] == pid, 'Cross-project group dependency')
            rank = integer(side.get('rank'), 'group rank', 1)
            hits = [r for r in g['rows'] if r['rank'] == rank]
            return hits[0]['entry_id'] if g['complete'] and len(hits) == 1 and not hits[0]['tied'] else None
        raise ValueError('Unknown source kind; represent byes as automatic advances, not matches')
    def match_result(mid):
        require(mid in matches, 'Unknown predecessor match')
        if mid in cache:
            return cache[mid]
        token = ('match', mid)
        require(token not in visiting, 'Cyclic match dependency')
        visiting.add(token)
        m = matches[mid]
        pid = m.get('project_id')
        require(pid in projects and len(m.get('sides', [])) == 2, 'Match needs project and two sources')
        rule = m.get('scoring', projects[pid].get('scoring_by_stage', {}).get(str(m.get('stage')), projects[pid]['scoring']))
        validate_rule(rule)
        ids = [resolve(s, pid) for s in m['sides']]
        require(ids[0] is None or ids[1] is None or ids[0] != ids[1], 'Entry plays itself')
        explicit = [match_result(i) for i in m.get('predecessors', [])]
        r = results.get(mid, {'status': 'pending'})
        status = r.get('status')
        require(status in STATUS, 'Unknown result status')
        require(not r.get('awarded_score') or status in ('walkover', 'retired'), 'Administrative score only belongs to an explicit abnormal result')
        out = {'match_id': mid, 'project_id': pid, 'group_id': m.get('group_id'), 'stage': m.get('stage'),
               'round': m.get('round'), 'entry_ids': ids, 'status': status,
               'winner_id': None, 'loser_id': None, 'score': deepcopy(r.get('score')),
               'reason': r.get('reason', ''), 'source': deepcopy(r.get('source')),
               'audit': deepcopy(r.get('audit')), 'statistics': None,
               'scoring_mode': rule['mode'], 'scoring_rule': deepcopy(rule), 'waived_predecessors': []}
        if status in FINAL:
            require(all(ids), f'{mid}: preceding result/rank unresolved')
            outcome_sources = {s['match_id'] for s in m['sides'] if s.get('kind') in ('winner', 'loser')}
            for predecessor in explicit:
                if predecessor['status'] in FINAL:
                    continue
                depid = predecessor['match_id']
                waiver = None
                if depid not in outcome_sources:
                    for side in m['sides']:
                        if side.get('kind') != 'group_rank':
                            continue
                        gid = side['group_id']
                        g, ranked = groups[gid], group_result(gid)
                        excluded = g.get('excluded_entries', {})
                        original = matches[depid]
                        original_ids = [s.get('entry_id') for s in original.get('sides', []) if s.get('kind') == 'entry']
                        affected = [i for i in original_ids if i in excluded]
                        if (ranked['complete'] and depid in g['match_ids'] and original.get('group_id') == gid
                                and original['project_id'] == pid and len(original_ids) == 2
                                and set(original_ids) <= set(g['entry_ids']) and affected):
                            waiver = {'match_id': depid, 'group_id': gid,
                                      'reason': '组内参赛项除名，相关组内赛果按明确裁定作废',
                                      'excluded_entries': {i: excluded[i] for i in affected}}
                            break
                require(waiver is not None, f'{mid}: predecessor not complete')
                out['waived_predecessors'].append(waiver)
            require(all(entries[i].get('kind') != 'placeholder' for i in ids), 'Cannot give a result to an unfilled reservation')
            require(r.get('winner_id') in ids, 'Winner must be one resolved match entry')
            if status == 'completed':
                winner, statistics = score_result(r.get('score'), rule)
                require(ids[winner] == r['winner_id'], 'Declared winner disagrees with score')
                out['statistics'] = statistics
            else:
                require(r.get('reason'), 'Walkover/retirement requires reason')
                if status == 'walkover':
                    require(not r.get('score'), 'Do not invent played scores for a walkover')
                elif r.get('score'):
                    _, out['statistics'] = score_result(r['score'], rule, complete=False)
                if r.get('awarded_score'):
                    require(rule['mode'] == 'points', 'Administrative score currently supports point games only')
                    require(r.get('awarded_score_reference'), 'Administrative score requires a confirmed ruling or report reference')
                    winner, statistics = score_result(r['awarded_score'], rule)
                    require(ids[winner] == r['winner_id'], 'Administrative score disagrees with declared winner')
                    if status == 'retired' and r.get('score'):
                        played = r['score'].get('games', [])
                        awarded = r['awarded_score']['games']
                        require(awarded[:len(played)] == played, 'Administrative score rewrites a completed played game')
                        partial = r['score'].get('partial_game')
                        if partial:
                            require(len(awarded) > len(played) and all(awarded[len(played)][i] >= partial[i] for i in range(2)),
                                    'Administrative score erases points already played')
                    out['awarded_score'] = deepcopy(r['awarded_score'])
                    out['awarded_score_reference'] = deepcopy(r['awarded_score_reference'])
                    out['awarded_statistics'] = statistics
            out['winner_id'] = r['winner_id']
            out['loser_id'] = ids[1 - ids.index(r['winner_id'])]
        else:
            require(not r.get('winner_id') and not r.get('score'), 'Pending/cancelled match cannot contain final score or winner')
            if status == 'cancelled':
                require(r.get('reason'), 'Cancellation requires reason')
        cache[mid] = out
        visiting.remove(token)
        return out
    for mid in matches:
        match_result(mid)
    for gid in groups:
        group_result(gid)
    placements, claims = [], {}
    for mid, m in matches.items():
        assignment = m.get('placement', {})
        if not assignment:
            continue
        require(set(assignment) <= {'winner', 'loser'}, 'Unknown placement assignment')
        for outcome, rank in assignment.items():
            integer(rank, 'placement rank', 1)
            eid = cache[mid].get(outcome + '_id')
            if not eid:
                continue
            pid = m['project_id']
            require((pid, rank) not in claims and (pid, eid) not in claims, 'Conflicting placement claims')
            claims[(pid, rank)] = claims[(pid, eid)] = mid
            placements.append({'project_id': pid, 'rank': rank, 'entry_id': eid, 'source_match_id': mid,
                               'source_outcome': outcome, 'status': '已核定'})
    for shared in data.get('shared_placements', []):
        pid, rank = shared.get('project_id'), integer(shared.get('rank'), 'shared rank', 1)
        mids = shared.get('source_match_ids', [])
        require(pid in projects and len(mids) >= 2 and len(mids) == len(set(mids)), 'Invalid shared placement sources')
        require(shared.get('outcome') in ('winner', 'loser') and shared.get('rule_reference'), 'Shared place requires explicit outcome and regulation')
        require(all(mid in matches and matches[mid]['project_id'] == pid for mid in mids), 'Foreign shared placement match')
        require((pid, rank) not in claims, 'Shared place conflicts with another assigned rank')
        outcome = shared['outcome']
        eids = [cache[mid].get(outcome + '_id') for mid in mids]
        if not all(eids):
            continue
        require(len(set(eids)) == len(eids), 'Duplicate entry in shared placement')
        require(all((pid, eid) not in claims for eid in eids), 'Entry has conflicting shared placement')
        claims[(pid, rank)] = mids
        for mid, eid in zip(mids, eids):
            claims[(pid, eid)] = mid
            placements.append({'project_id': pid, 'rank': rank, 'entry_id': eid, 'source_match_id': mid,
                               'source_outcome': outcome, 'status': '并列名次（规程已确认）'})
    annulled = {mid for g in groups.values() for mid in g['match_ids']
                if any(i in g.get('excluded_entries', {}) for i in cache[mid]['entry_ids'])}
    return {'schema': 'gamemaster.results.v1', 'event_id': data.get('event_id'),
            'event_name': data.get('event_name', ''), 'synthetic': data.get('synthetic', False),
            'input_sha256': sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
            'projects': deepcopy(list(projects.values())), 'entries': deepcopy(list(entries.values())),
            'athletes': deepcopy(data.get('athletes', [])), 'matches': list(cache.values()),
            'groups': list(gcache.values()), 'placements': sorted(placements, key=lambda r: (r['project_id'], r['rank'])),
            'pending_match_ids': [i for i, r in cache.items() if r['status'] == 'pending' and i not in annulled],
            'cancelled_match_ids': [i for i, r in cache.items() if r['status'] == 'cancelled'],
            'annulled_match_ids': sorted(annulled),
            'unresolved_match_ids': [i for i, r in cache.items() if r['status'] not in FINAL and i not in annulled]}


def entry_text(entry, people):
    members = [people[i] for i in entry.get('member_ids', []) if i in people]
    label = entry.get('label') or '／'.join(m.get('name', m['id']) for m in members) or entry['id']
    units = list(dict.fromkeys(m.get('unit', m.get('club', '')) for m in members))
    return entry.get('club') or '／'.join(filter(None, units)), label


def export_report(report, destination, xlsx=True):
    """Export immutable result snapshots, not an editable spreadsheet calculation model."""
    destination = Path(destination)
    require(not destination.exists(), 'Destination already exists; create a new version directory')
    if xlsx:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter
    destination.mkdir(parents=True)
    (destination / '成绩数据.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    entries = index(report['entries'])
    people = index(report.get('athletes', []))
    names = {i: entry_text(e, people) for i, e in entries.items()}
    def safe(v):
        return "'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v
    summary = [['项目', '组号', '名次', '队名', '姓名', '备注', '报名项ID', '依据场次ID']]
    sheets = []
    for p in report['projects']:
        rows = [[p.get('name', p['id'])], [], ['组号', '名次', '队名', '姓名', '备注']]
        for g in report['groups']:
            if g['project_id'] != p['id']:
                continue
            for r in g['rows']:
                team, name = names[r['entry_id']]
                note = r['status'] + ('；并列待裁定' if r['tied'] else '')
                rows.append([g['id'], r['rank'], team, name, note])
                summary.append([p.get('name', p['id']), g['id'], r['rank'], team, name, note, r['entry_id'], ''])
        if len(rows) > 3:
            sheets.append((p['id'] + '_循环成绩', rows))
        final = [[p.get('name', p['id']) + ' 淘汰名次'], [], ['名次', '队名', '姓名', '备注']]
        for r in report['placements']:
            if r['project_id'] != p['id']:
                continue
            team, name = names[r['entry_id']]
            final.append([r['rank'], team, name, r['status']])
            summary.append([p.get('name', p['id']), '', r['rank'], team, name, r['status'], r['entry_id'], r['source_match_id']])
        if len(final) > 3:
            sheets.append((p['id'] + '_淘汰成绩', final))
    with (destination / '成绩排名.csv').open('w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f).writerows([[safe(v) for v in row] for row in summary])
    detail = [['场次ID', '项目ID', '阶段', '组号', '轮次', '甲方', '乙方', '状态', '胜方', '完整结构化比分', '原因', '行政判定比分（非实际对局）', '行政判定依据']]
    for m in report['matches']:
        a, b = [names.get(i, ('', '待确定'))[1] for i in m['entry_ids']]
        detail.append([m['match_id'], m['project_id'], m['stage'], m['group_id'], m['round'], a, b,
                       STATUS[m['status']], names.get(m['winner_id'], ('', ''))[1],
                       json.dumps(m['score'], ensure_ascii=False) if m['score'] else '', m['reason'],
                       json.dumps(m['awarded_score'], ensure_ascii=False) if m.get('awarded_score') else '',
                       json.dumps(m.get('awarded_score_reference'), ensure_ascii=False) if m.get('awarded_score_reference') else ''])
    with (destination / '逐场成绩.csv').open('w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f).writerows([[safe(v) for v in row] for row in detail])
    if xlsx:
        wb = Workbook()
        wb.remove(wb.active)
        sheets += [('逐场成绩', detail), ('导出说明', [['字段', '说明'], ['数据类型', '核验后静态成绩快照；修改比分后重新运行引擎'],
            ['样例性质', '合成测试，不代表真实成绩' if report['synthetic'] else '来源见输入数据与校验记录'],
            ['未决场次', len(report.get('unresolved_match_ids', report['pending_match_ids']))],
            ['取消场次', len(report.get('cancelled_match_ids', []))],
            ['明确作废的组内场次', len(report.get('annulled_match_ids', []))], ['来源校验值', report['input_sha256']],
            ['网球统计', 'games表示盘；points表示普通局，抢十不计普通局；完整比分保留结构']])]
        used = set()
        for n, (title, rows) in enumerate(sheets, 1):
            title = ''.join(c if c not in '[]:*?/\\' else '_' for c in title)[:26]
            title = f'{n}_{title}'[:31]
            require(title not in used, 'Duplicate sheet title')
            used.add(title)
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append(row)
            header = 3 if rows and len(rows) > 2 and not rows[1] else 1
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.font = Font(name='Arial', size=11, bold=cell.row in (1, header))
                    cell.alignment = Alignment(vertical='center', wrap_text=True)
                    cell.border = Border(bottom=Side(style='hair', color='BFBFBF'))
            ws.freeze_panes = 'A' + str(header + 1)
            for c in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(c)].width = 20 if c < 6 else 30
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = f'1:{header}'
        path = destination / '成绩排名.xlsx'
        wb.save(path)
        reopened = load_workbook(path)
        require(len(reopened.sheetnames) == len(sheets), 'Workbook read-back mismatch')
        require(not any(c.data_type == 'f' for ws in reopened for row in ws for c in row), 'Unexpected formula in static report')
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path, help='Match graph with project scoring and results')
    parser.add_argument('--output', required=True, type=Path, help='New export directory')
    parser.add_argument('--no-xlsx', action='store_true', help='Standard-library-only JSON/CSV')
    args = parser.parse_args()
    report = evaluate(json.loads(args.input.read_text(encoding='utf-8')))
    export_report(report, args.output, not args.no_xlsx)
    print(json.dumps({'matches': len(report['matches']), 'placements': len(report['placements']),
                      'pending': len(report['pending_match_ids']), 'output': str(args.output)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
