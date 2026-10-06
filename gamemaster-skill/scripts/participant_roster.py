"""Merge reviewed registration changes and export the bundled roster layout."""
import argparse
from collections import Counter
from copy import deepcopy, copy
import json
from pathlib import Path
import re
from openpyxl import load_workbook


def need(ok, message):
    if not ok:
        raise ValueError(message)


def text(value):
    return isinstance(value, str) and bool(value.strip())


def merge(base, batch):
    out = deepcopy(base)
    entries = {e['id']: e for e in out['entries']}
    need(len(entries) == len(out['entries']), '基线报名ID重复')
    history = out.setdefault('history', [])
    applied = {h['operation']['id']: h['operation'] for h in history}
    changed = False
    for op in batch['operations']:
        need(text(op.get('id')) and text(op.get('source')), '操作须有唯一ID与来源')
        if op['id'] in applied:
            need(applied[op['id']] == op, '相同操作ID内容冲突')
            continue
        action = op['action']
        need(action in ('add', 'replace', 'withdraw'), '不支持的操作')
        entry = deepcopy(op.get('entry'))
        key = op.get('entry_id') if action == 'withdraw' else entry['id']
        before = deepcopy(entries.get(key))
        if action == 'add':
            need(before is None or before == entry, '新增ID已存在且内容冲突，请明确更正')
        else:
            need(before is not None, '更正或退赛目标不存在')
        if action == 'withdraw':
            entry = deepcopy(before)
            entry['status'] = 'withdrawn'
        need(text(entry.get('id')) and entry.get('status') in ('active', 'withdrawn'), '报名ID或状态无效')
        need(text(entry.get('source')), '报名记录须有来源')
        entries[key] = entry
        history.append({'operation': deepcopy(op), 'before': before, 'after': deepcopy(entry)})
        applied[op['id']] = deepcopy(op)
        changed |= before != entry
    out['entries'] = list(entries.values())
    if changed:
        out['reviewed'] = False
        out['version'] = out.get('version', 1) + 1
    return out


def validate(data, system=False):
    need(data.get('reviewed') is True and data.get('unresolved') == [], '先完成审核并解决未决项')
    need(type(data.get('exclusive')) is bool, '须依据规程明确是否禁止兼项exclusive')
    projects = data.get('projects', [])
    need(bool(projects), '未定义项目')
    index = {}
    system_ids = set()
    for p in projects:
        need(text(p.get('key')) and text(p.get('name')), '项目编号/名称缺失')
        need(p['key'] not in index, '项目编号重复')
        need(p.get('type') in ('MS','WS','SS','MD','WD','XD','SD'), '该项目需专用模板')
        if 'expected_count' in p:
            need(type(p['expected_count']) is int and p['expected_count'] >= 0, '项目人数无效')
        if system:
            need(text(data.get('event_id')) and text(p.get('system_id')), '系统导入须提供目标赛事ID及项目XMID')
            need(p['system_id'] not in system_ids, '系统项目ID重复')
            system_ids.add(p['system_id'])
        index[p['key']] = p
    ids, per_project, global_people, seeds, technical = set(), set(), {}, set(), set()
    people = {}
    counts = Counter()
    real_counts = Counter()
    active = []
    for e in data.get('entries', []):
        need(text(e.get('id')) and e['id'] not in ids, '报名ID缺失或重复')
        ids.add(e['id'])
        need(e.get('status') in ('active','withdrawn'), '报名状态无效')
        need(e.get('project') in index, '报名项目未定义')
        need(text(e.get('source')), '报名来源缺失')
        if e['status'] == 'withdrawn':
            continue
        p = index[e['project']]
        need(e.get('kind', 'real') in ('real', 'placeholder'), '报名类型无效')
        if e.get('kind') == 'placeholder':
            need(not system, '预留名额尚未验证跑兔导入；先替换为真实报名项，不能伪造性别或身份')
            need(e.get('members') == [] and text(e.get('label')), '预留名额须有标签且不带虚构运动员')
            need(e.get('seed') in (None, '', 0, '0'), '预留名额不能指定运动员种子')
            policy = e.get('reservation', {})
            need(policy.get('confirmed') is True and all(text(policy.get(k)) for k in
                 ('replacement_deadline', 'eligibility', 'unfilled', 'allocation')), '预留名额须确认替换期限、资格、空位处理和分配办法')
            counts[p['key']] += 1
            active.append(e)
            continue
        n = 1 if p['type'].endswith('S') else 2
        members = e.get('members', [])
        need(len(members) == n, '单打须1人、双打须2人，缺搭档不能导出')
        genders = []
        for member in members:
            for k in ('id','name','unit'):
                need(text(member.get(k)), f'运动员{k}缺失')
            mid = member['id']
            need((p['key'], mid) not in per_project, '同一运动员在同项目重复报名或双打成员重复')
            per_project.add((p['key'],mid))
            if data['exclusive']:
                need(mid not in global_people, '规程禁止兼项，发现跨项目重复运动员')
            global_people[mid] = True
            if mid in people:
                need(people[mid] == member, '同一运动员ID的信息冲突，先统一资料')
            people[mid] = member
            gender = member.get('gender', '')
            need(gender in ('M','W') if system else gender in ('','M','W'),
                 '系统导入性别必填M/W；缺失须核对原件或向用户确认')
            genders.append(gender)
            for k in ('identity','phone'):
                need(k not in member or isinstance(member[k], str), f'{k}必须以文本保留，避免号码精度丢失')
            need('age' not in member or type(member['age']) is int and member['age'] >= 0, '年龄须为非负整数')
        if p['type'] in ('MS','MD'):
            need(all(g == 'M' for g in genders), '男子项目性别缺失或不符')
        if p['type'] in ('WS','WD'):
            need(all(g == 'W' for g in genders), '女子项目性别缺失或不符')
        if p['type'] == 'XD':
            need(sorted(genders) == ['M','W'], '混双须确认一男一女')
        for k, seen in [('seed', seeds), ('technical_no', technical)]:
            value = e.get(k)
            if k == 'seed' and value in (0, '0') and not isinstance(value, bool):
                continue  # Zero denotes an unseeded entry, not a repeated seed rank.
            if value not in (None, ''):
                need(isinstance(value, (str,int)) and not isinstance(value,bool), f'{k}类型无效')
                marker = (p['key'], str(value)) if k == 'seed' else (p['key'], tuple(m['unit'] for m in members), str(value))
                need(marker not in seen, f'{k}重复')
                seen.add(marker)
        counts[p['key']] += 1
        real_counts[p['key']] += 1
        active.append(e)
    for p in projects:
        if 'expected_real_count' in p:
            need(type(p['expected_real_count']) is int and p['expected_real_count'] >= 0,
                 '真实报名项数量无效')
            need(real_counts[p['key']] == p['expected_real_count'], f"{p['name']}真实报名项数量不符")
        if 'expected_count' in p:
            need(counts[p['key']] == p['expected_count'], f"{p['name']}名单数量与方案不符：{counts[p['key']]} / {p['expected_count']}")
    return active, counts


def summarize(data):
    """Separate real entries, reserved entry slots, and unique known athletes."""
    active, counts = validate(data)
    result = {}
    for p in data['projects']:
        rows = [e for e in active if e['project'] == p['key']]
        reserved = sum(e.get('kind') == 'placeholder' for e in rows)
        result[p['key']] = {'planned_entries': counts[p['key']], 'reserved_entries': reserved,
                            'real_entries': counts[p['key']] - reserved,
                            'known_athletes': len({m['id'] for e in rows for m in e['members']})}
    return result


def export(data, destination, system=False):
    active, counts = validate(data, system)
    target = Path(destination)
    need(target.suffix.lower() == '.xlsx' and not target.exists(), '输出须为尚不存在的.xlsx文件')
    book = load_workbook(Path(__file__).resolve().parents[1] / 'assets/participant-roster.xlsx')
    prototypes = list(book)
    for number, p in enumerate(data['projects'], 1):
        single = p['type'].endswith('S')
        sheet = book.copy_worksheet(prototypes[0 if single else 1])
        title = re.sub(r'[\\/*?:\[\]]', '-', p['name'])
        sheet.title = f'{title[:20]}_{p["type"]}_{number}'
        sheet['A1'] = f'({p["name"]}){p["name"]} --- {p["type"]}'
        sheet['A1'].data_type = 's'
        sheet['A2'] = p.get('system_id') if system else None
        if system:
            sheet['A2'].data_type = 's'
        sheet['B2'] = p['name']; sheet['B2'].data_type = 's'
        sheet['C2'] = p['type']
        styles = [copy(c._style) for c in sheet[4]]
        selected = [e for e in active if e['project'] == p['key']]
        for rownum, e in enumerate(selected, 4):
            sheet.row_dimensions[rownum].height = sheet.row_dimensions[4].height
            if e.get('kind') == 'placeholder':
                label = e['label']
                if single:
                    values = ['预留名额', label, None, None, None, None, None, None]
                else:
                    values = ['预留名额', label + '-A位', None, None, None, None,
                              '预留名额', label + '-B位', None, None, None, None] + [None] * 8
            elif single:
                a = e['members'][0]
                values = [a['unit'],a['name'],a.get('gender'),e.get('technical_no'),e.get('seed'),a.get('identity'),a.get('age'),a.get('phone')]
            else:
                values = []
                for member in e['members']:
                    values.extend(member.get(k) for k in ('unit','name','identity','age','gender','phone'))
                values.extend([None]*6 + [e.get('technical_no'), e.get('seed')])
            for column, value in enumerate(values,1):
                c = sheet.cell(rownum,column,value)
                c._style = copy(styles[column-1])
                if isinstance(value,str):
                    c.data_type='s';c.number_format='@'
    for sheet in prototypes:
        book.remove(sheet)
    target.parent.mkdir(parents=True,exist_ok=True)
    book.save(target)
    check = load_workbook(target)
    need(len(check.worksheets) == len(data['projects']), '输出工作表数不符')
    for sheet,p in zip(check.worksheets,data['projects']):
        actual = sum(sheet.cell(r,2).value is not None for r in range(4,sheet.max_row+1))
        need(actual == counts[p['key']], '输出行数不符')
    return dict(counts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command',required=True)
    m = sub.add_parser('merge');m.add_argument('base');m.add_argument('batch');m.add_argument('output')
    e = sub.add_parser('export');e.add_argument('input');e.add_argument('output');e.add_argument('--system',action='store_true')
    args = parser.parse_args()
    read = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
    try:
        if args.command == 'merge':
            need(not Path(args.output).exists(), '输出已存在，请保留原版本')
            result = merge(read(args.base),read(args.batch))
            Path(args.output).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
            print(args.output)
        else:
            print(json.dumps(export(read(args.input),args.output,args.system),ensure_ascii=False))
    except (ValueError, KeyError, TypeError) as exc:
        parser.exit(2, str(exc)+'\n')


if __name__ == '__main__':
    main()
