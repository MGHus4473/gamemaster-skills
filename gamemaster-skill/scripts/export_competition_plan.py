"""Export confirmed, conventional plans to the bundled 21-column XLSX template."""
import argparse
import json
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

KEYS = 'id group type title stage extra entrants groups advance start_rank previous_groups previous_start previous_end grab games points cap scoring format rotation draw'.split()
INTS = 'stage extra entrants groups advance start_rank games points cap scoring rotation'.split()
TYPES = {'MS', 'WS', 'MD', 'WD', 'XD', 'SS', 'SD', 'MT', 'WT', 'XT'}
DRAWS = {'0', 'ABLZJC', 'AB', 'AC', 'AD', 'ABT', 'ABS', 'ALLSJ', 'ABTZTDWSLBXY'}


def require(ok, message):
    if not ok:
        raise ValueError(message)


def validate(data):
    require(data.get('confirmed') is True, '先补齐并确认关键参数，再生成正式表')
    rows = data.get('rows')
    require(isinstance(rows, list) and rows, 'rows不能为空')
    indexed = {}
    for i, r in enumerate(rows, 1):
        prefix = f'第{i}行：'
        for key in KEYS:
            if key not in {'previous_groups', 'previous_start', 'previous_end', 'grab'}:
                require(key in r and r[key] is not None, prefix + f'缺少{key}')
        for key in INTS:
            require(type(r[key]) is int and r[key] >= 0, prefix + f'{key}须为非负整数')
        for key in ['id', 'group', 'title', 'type', 'format', 'draw']:
            require(isinstance(r[key], str) and r[key].strip(), prefix + f'{key}须为非空文本')
        require(r['type'] in TYPES, prefix + '项目类型需单独核验')
        require(r['stage'] in (1, 2) and r['extra'] == 1, prefix + '脚本仅支持1/2阶段、附加1')
        require(r['entrants'] >= 2 and r['start_rank'] >= 1, prefix + '人数或起始名次无效')
        require(r['games'] in (1, 3, 5, 7, 9, 11), prefix + '局数不支持')
        require(0 < r['points'] <= r['cap'] and r['scoring'] in (1, 2), prefix + '计分参数无效')
        require(0 < r['advance'] <= r['entrants'], prefix + '循环晋级数或淘汰取名次数量无效')
        key = (r['id'], r['stage'])
        require(key not in indexed, prefix + '项目阶段重复')
        indexed[key] = r
        if r['format'] == 'XH':
            sizes = r.get('group_sizes')
            require(isinstance(sizes, list) and len(sizes) == r['groups'] and sizes,
                    prefix + '须提供各组人数group_sizes')
            require(all(type(n) is int and n >= 2 for n in sizes), prefix + '小组人数须至少2')
            require(sum(sizes) == r['entrants'] and r['advance'] <= min(sizes), prefix + '分组人数或晋级数矛盾')
            require(r['rotation'] in (0, 1, 2, 3) and r['draw'] in ('0', '1'), prefix + '循环轮转/抽签未支持')
            require(r['stage'] == 1 or r['groups'] == 1, prefix + '模板要求二阶段循环组数为1')
        elif r['format'] == 'TT':
            require(r['groups'] == 0 and r['rotation'] == 0, prefix + '淘汰组数及轮转应为0')
            require(r['draw'] in (DRAWS if r['stage'] == 2 else {'0'}), prefix + '抽签模式未支持')
        else:
            raise ValueError(prefix + '只支持XH/TT')
        require(r.get('grab') in (None, ''), prefix + '是否抢号应留空')
        if r['stage'] == 1:
            require(all(r.get(k, 0) == 0 for k in KEYS[10:13]), prefix + '首阶段来源应为0')
    for r in rows:
        if r['stage'] == 1:
            continue
        prev = indexed.get((r['id'], 1))
        require(prev is not None and prev['format'] == 'XH', '第二阶段须有同项目第一阶段循环来源')
        require((prev['group'], prev['type']) == (r['group'], r['type']), '阶段间组别/类型不一致')
        for k in KEYS[10:13]:
            require(type(r.get(k)) is int and r[k] > 0, f'第二阶段缺少正整数{k}')
        require(r['previous_groups'] == prev['groups'], '前阶段组数不一致')
        require(1 == r['previous_start'] <= r['previous_end'] == prev['advance'],
                '脚本仅支持各组从第1名到晋级名次的完整主线')
        require(r['entrants'] == r['previous_groups'] * r['previous_end'], '下一阶段人数与来源席位不符')
    return rows


def export(data, destination):
    rows = validate(data)
    destination = Path(destination)
    template = Path(__file__).resolve().parents[1] / 'assets/competition-plan.xlsx'
    require(destination.suffix.lower() == '.xlsx', '输出须为.xlsx')
    require(not destination.exists(), '目标文件已存在，请另选文件名')
    book = load_workbook(template)
    sheet = book.active
    styles = [copy(c._style) for c in sheet[3]]
    height = sheet.row_dimensions[3].height
    for number, row in enumerate(rows, 3):
        sheet.row_dimensions[number].height = height
        for col, key in enumerate(KEYS, 1):
            value = row.get(key, 0 if key.startswith('previous_') else None)
            cell = sheet.cell(number, col, value)
            cell._style = copy(styles[col - 1])
            if isinstance(value, str):
                cell.data_type = 's'
            if key == 'id':
                cell.number_format = '@'
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    check = load_workbook(destination)
    assert check.active.max_column == 21
    assert check.active.max_row == len(rows) + 2
    for number, row in enumerate(rows, 3):
        assert check.active.cell(number, 1).value == row['id']
    return destination


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    try:
        result = export(json.loads(args.input.read_text(encoding='utf-8')), args.output)
        print(result)
    except (ValueError, KeyError) as exc:
        parser.exit(2, str(exc) + '\n')
