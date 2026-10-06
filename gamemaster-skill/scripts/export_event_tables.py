#!/usr/bin/env python3
"""Offline draw, schedule and extended multi-sport plan workbooks; no platform IDs required."""
import argparse
from copy import copy
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from draw_engine import digest, validate_project
from schedule_engine import validate as validate_schedule, slots_for
from export_competition_plan import KEYS
from results_engine import validate_rule


def require(ok, message):
    if not ok:
        raise ValueError(message)


def save(book, output):
    path = Path(output)
    require(path.suffix.lower() == '.xlsx' and not path.exists(), 'Use a new .xlsx output path')
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Imported names and labels are data, never Excel formulas.
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    check = load_workbook(path, data_only=False)
    require(check.sheetnames == book.sheetnames, 'Workbook sheet mismatch')
    for original, restored in zip(book, check):
        require(original.max_row == restored.max_row and original.max_column == restored.max_column,
                'Workbook dimensions mismatch')
        for row in original:
            for cell in row:
                expected = None if cell.value == '' else cell.value
                require(restored[cell.coordinate].value == expected, 'Workbook value changed: ' + cell.coordinate)
    return {'output': str(path), 'sheets': book.sheetnames, 'readback': 'passed', 'mode': 'offline'}


def table(book, name, headers, rows, header_row=1):
    sheet = book.create_sheet(name)
    if header_row == 2:
        sheet.append(['本地赛事文件；标识为本地ID，平台导入另行适配'])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[header_row]:
        cell.fill = PatternFill('solid', fgColor='24476A')
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for col in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 24
    sheet.freeze_panes = f'A{header_row + 1}'
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}'
    sheet.print_title_rows = f'{header_row}:{header_row}'
    return sheet


def book_without_default():
    book = Workbook(); book.remove(book.active)
    return book


def export_draw(config, result, output):
    require(result['event_id'] == config['event_id'] and result['input_sha256'] == digest(config), 'Draw input version mismatch')
    projects = {p['id']: p for p in config['projects']}
    require(len(projects) == len(config['projects']) == len(result['projects'])
            and set(projects) == {p['project_id'] for p in result['projects']}, 'Draw project mismatch')
    require(result['random_seed'] == config['random_seed'], 'Draw audit random seed mismatch')
    headers = ['项目ID', '阶段', '附加', '组别', '项目名称', '项目类型', '赛事种类', '组号/轮次号',
               '位置号', '2分区', '4分区', '8分区', '16分区', '队伍名称', '姓名', '性别',
               '队内技术号', '种子号', '前阶段名次', '识别码']
    rows = []
    for project in result['projects']:
        source = projects[project['project_id']]
        validate_project(source, project)
        assignments = list(project['assignments'])
        assignments += [{'entry_id': '', 'name': '轮空', 'club': '', 'seed': 0, 'group': 1, 'position': pos}
                        for pos in project.get('byes', [])]
        for row in sorted(assignments, key=lambda a: (a['group'], a['position'])):
            rows.append([source['id'], 1, 1, source.get('name', source['id']), source.get('name', source['id']),
                         source.get('type'), source.get('sport'), row['group'], row['position'],
                         None, None, None, None, row.get('club'), row.get('name'), None, None,
                         row.get('seed') or None, None, row['entry_id']])
    book = book_without_default()
    table(book, '抽签工作表', headers, rows, header_row=2)
    table(book, '抽签记录', ['项目', '内容'], [
        ['赛事本地ID', config['event_id']], ['随机种子', result['random_seed']],
        ['输入SHA256', result['input_sha256']], ['模式', '离线；预留身份未知，单位回避仅覆盖已知单位']])
    return save(book, output)


def side_label(side):
    if side.get('label'):
        return side['label']
    kind = side.get('kind')
    if kind == 'entry':
        return side['entry_id']
    if kind in ('winner', 'loser'):
        return side['match_id'] + ('胜者' if kind == 'winner' else '负者')
    if kind == 'group_rank':
        return side['group_id'] + '第' + str(side['rank']) + '名'
    return '待定'


def export_schedule(data, result, output):
    require(result['event_id'] == data['event_id'], 'Schedule event mismatch')
    expected_hash = hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    require(result['input_sha256'] == expected_hash, 'Schedule input version mismatch')
    require(result['slots'] == slots_for(data), 'Schedule grid differs from confirmed sessions')
    report = validate_schedule(data, result)
    require(report['complete'], 'Incomplete or invalid schedule: ' + str(report))
    matches = {m['id']: m for m in data['matches']}
    assignments = {(a['slot_index'], a['court']): a for a in result['assignments']}
    rows, detail = [], []
    for slot in result['slots']:
        row = [slot['start'][:10], slot['start'][11:], str(slot['scene'])]
        for court in data['courts']:
            rec = assignments.get((slot['index'], court))
            if not rec:
                row.append(None); continue
            m = matches[rec['match_id']]
            labels = [side_label(s) for s in m.get('sides', [])]
            labels += ['待定'] * (2 - len(labels))
            row.append('\n'.join([m.get('code', m['id']), m.get('project_name', m['project_id']),
                                  labels[0] + ' VS ' + labels[1], '本地ID：' + m['id']]))
            detail.append([m['id'], m.get('code'), m.get('project_name', m['project_id']), m.get('stage'),
                           m.get('round'), rec['start'][:10], rec['start'][11:], rec['end'][11:],
                           rec['scene'], rec['section'], rec['court'], *labels[:2],
                           '、'.join(m.get('predecessors', []))])
        rows.append(row)
    book = book_without_default()
    grid = table(book, '赛事编排工作表', ['日期', '时间', '场序'] + [f'第{c}号场地' for c in data['courts']], rows)
    grid.freeze_panes = 'D2'
    for col in range(4, grid.max_column + 1):
        grid.column_dimensions[get_column_letter(col)].width = 38
    for row in range(2, grid.max_row + 1):
        grid.row_dimensions[row].height = 100
    table(book, '场次工作表', ['本地场次ID', '赛号', '项目', '阶段', '轮次', '日期', '开始', '预计结束',
                                '场序', '小节', '场地', '对阵一', '对阵二', '前置场次'], detail)
    table(book, '编排核验', ['事项', '内容'], [
        ['赛事本地ID', data['event_id']], ['模式', '离线排期；未绑定平台ID'],
        ['参数', json.dumps({k: v for k, v in data.items() if k != 'matches'}, ensure_ascii=False)],
        ['校验', json.dumps(report, ensure_ascii=False)],
        ['候选口径', data['conflict_scope']], ['时间口径', '计划时间，实际超时须更新休息及后续场次']])
    return save(book, output)


def export_plan(data, output):
    """Preserve the 21-column outline plus full sport rules, never fake tennis point columns."""
    require(data.get('confirmed') is True, 'Confirm plan first')
    rows = data.get('rows', [])
    require(rows and all(r.get('id') and r.get('sport') and isinstance(r.get('rules'), dict) and r['rules'] for r in rows), 'Each row needs id, sport and full rules')
    seen = set()
    for row in rows:
        marker = (row['id'], row.get('stage', 1), row.get('extra', 1))
        require(marker not in seen, 'Duplicate plan stage')
        seen.add(marker)
        require(row['sport'] in ('badminton', 'table_tennis', 'tennis', 'pickleball'), 'Unknown sport')
        validate_rule(row['rules'])
        require(row['rules']['mode'] == ('tennis' if row['sport'] == 'tennis' else 'points'), 'Scoring mode does not match sport')
        require(all(isinstance(row.get(k), str) and row[k].strip() for k in ('id', 'group', 'title')), 'Plan labels required')
        require(type(row.get('stage')) is int and row['stage'] in (1, 2) and row.get('extra') == 1, 'Only stages 1/2 and ordinary branch 1 supported')
        require(type(row.get('start_rank')) is int and row['start_rank'] >= 1, 'Starting rank required')
        require(type(row.get('groups')) is int and row['groups'] >= 0, 'Group count required')
        require(type(row.get('entrants')) is int and row['entrants'] >= 2, 'Invalid entry count')
        require(row.get('format') in ('XH', 'TT') and row.get('type') in ('MS', 'WS', 'SS', 'MD', 'WD', 'XD', 'SD'), 'Unsupported plan shape')
        if row['format'] == 'XH':
            sizes = row.get('group_sizes', [])
            require(sizes and all(type(n) is int and n >= 2 for n in sizes) and sum(sizes) == row['entrants'], 'Group capacity mismatch')
            require(row.get('groups') == len(sizes), 'Group count mismatch')
        require(type(row.get('advance')) is int and 1 <= row['advance'] <= row['entrants'], 'Confirm advancing/placement count')
        if row['format'] == 'XH':
            require(row['advance'] <= min(row['group_sizes']), 'Advancement exceeds smallest group')
        else:
            require(row['groups'] == 0, 'Knockout group count must be zero')
    for row in rows:
        if row['stage'] == 2:
            previous = next((r for r in rows if r['id'] == row['id'] and r['stage'] == 1), None)
            require(previous and previous['format'] == 'XH', 'Stage 2 needs preceding group stage')
            require(row.get('previous_groups') == previous['groups']
                    and row.get('previous_start') == 1 and row.get('previous_end') == previous['advance'], 'Advancement source mismatch')
            require(row['entrants'] == previous['groups'] * previous['advance'], 'Stage capacity mismatch')
            require((row['sport'], row['type']) == (previous['sport'], previous['type']), 'Sport/type differs between stages')
    book = load_workbook(Path(__file__).resolve().parents[1] / 'assets/competition-plan.xlsx')
    sheet = book.active
    sheet.title = '竞赛方案'
    for number, original in enumerate(rows, 3):
        row = dict(original)
        if row['sport'] == 'tennis':
            for key in ('games', 'points', 'cap'):
                row[key] = None
        else:
            row.update(games=row['rules']['best_of'], points=row['rules']['target'], cap=row['rules'].get('cap'))
        for column, key in enumerate(KEYS, 1):
            cell = sheet.cell(number, column)
            cell._style = copy(sheet.cell(3, column)._style)
            cell.value = row.get(key)
    table(book, '运动规则', ['项目本地ID', '阶段', '运动', '完整规则JSON'],
          [[r['id'], r.get('stage', 1), r['sport'], json.dumps(r['rules'], ensure_ascii=False, sort_keys=True)] for r in rows])
    table(book, '文件说明', ['事项', '内容'], [['模式', '扩展离线方案；不直接作为21列平台导入文件'],
        ['网球', '盘/局/抢七结构在运动规则表完整保留，单一局分字段留空'],
        ['核验范围', '文件容量与结构；实际规则、阶段衔接、对阵及排期另由对应引擎核验']])
    return save(book, output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    for command in ('draw', 'schedule'):
        p = sub.add_parser(command)
        p.add_argument('input'); p.add_argument('result'); p.add_argument('output')
    p = sub.add_parser('plan'); p.add_argument('input'); p.add_argument('output')
    a = parser.parse_args()
    read = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
    try:
        if a.command == 'plan':
            report = export_plan(read(a.input), a.output)
        else:
            report = {'draw': export_draw, 'schedule': export_schedule}[a.command](read(a.input), read(a.result), a.output)
        print(json.dumps(report, ensure_ascii=False))
    except (ValueError, KeyError, TypeError, OSError) as exc:
        parser.exit(2, str(exc) + '\n')


if __name__ == '__main__':
    main()
