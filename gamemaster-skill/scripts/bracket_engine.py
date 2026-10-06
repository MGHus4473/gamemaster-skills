#!/usr/bin/env python3
"""Offline match graph: confirmed draw -> round robin / knockout / group knockout.

Only Excel export needs openpyxl. Generation uses the Python standard library.
"""
import argparse
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5


SCHEMA = 'gamemaster.brackets.v1'
HEADERS17 = ['序号', '赛事名称', '项目ID', '组别', '项目', '组号', '轮次', '阶段',
             '附加', '淘汰赛名次', '队伍名称1', '人员姓名1', '队伍名称2', '人员姓名2',
             '场次号（勿动）', '原赛号', '新赛号']


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(value):
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                             separators=(',', ':')).encode()).hexdigest()


def stable_id(kind, *parts):
    return 'LOCAL-' + kind.upper() + '-' + str(uuid5(
        NAMESPACE_URL, 'urn:gamemaster:bracket:v1:' + json.dumps(
            [kind, *parts], ensure_ascii=False, separators=(',', ':'))))


def generate(config):
    """Generate a fresh graph without mutating the source or resolving results."""
    require(config.get('reviewed') is True, 'Confirm roster, format and draw before generating')
    require(isinstance(config.get('event_id'), str) and config['event_id'], 'event_id required')
    projects = config.get('projects', [])
    entries = deepcopy(config.get('entries', []))
    athletes = deepcopy(config.get('athletes', []))
    require(projects and len({p['id'] for p in projects}) == len(projects), 'Project IDs must be unique')
    require(len({e['id'] for e in entries}) == len(entries), 'Entry IDs must be unique across projects')
    require(len({a['id'] for a in athletes}) == len(athletes), 'Athlete IDs must be unique')
    require(all(isinstance(x.get('id'), str) and x['id'] for x in projects + entries + athletes), 'IDs must be non-empty strings')
    people = {a['id']: a for a in athletes}
    by_entry = {e['id']: e for e in entries}
    project_ids = {p['id'] for p in projects}
    require(all(e.get('project_id') in project_ids for e in entries), 'Entry references unknown project')
    for e in entries:
        members = e.get('member_ids', [])
        require(isinstance(members, list) and len(members) == len(set(members)), 'Duplicate member within entry')
        if e.get('kind') == 'placeholder':
            require(not members, 'Unfilled placeholder must not contain invented athlete identities')
        else:
            require(members and set(members) <= people.keys(), 'Real entries require known member_ids')
        e.setdefault('label', '／'.join(people[i].get('name', i) for i in members) or '预留名额')
    draw = config.get('draw', {})
    if draw.get('event_id'):
        require(draw['event_id'] == config['event_id'], 'Draw belongs to a different event')
    draw_projects = draw.get('projects', [])
    require(len({p['project_id'] for p in draw_projects}) == len(draw_projects), 'Duplicate draw project')
    require({p['project_id'] for p in draw_projects} <= project_ids, 'Unknown draw project')
    draws = {p['project_id']: p for p in draw_projects}
    out = {'schema': SCHEMA, 'event_id': config['event_id'],
           'event_name': config.get('event_name', config['event_id']),
           'input_sha256': digest(config), 'athletes': athletes, 'entries': entries,
           'projects': [], 'groups': [], 'matches': [], 'automatic_advances': [],
           'outcome_disjoint_pairs': [], 'platform_event_id': None,
           'unresolved_identity_entries': [e['id'] for e in entries if e.get('kind') == 'placeholder']}

    for original in projects:
        p = deepcopy(original)
        pid = p['id']
        p.setdefault('name', pid)
        require(p.get('sport') in ('badminton', 'table_tennis', 'tennis', 'pickleball'), 'Explicit supported sport required')
        require(p.get('format') in ('knockout', 'round_robin', 'groups_knockout'), 'Unsupported format')
        require(p.get('format') != 'round_robin' or 'classification_places' not in p,
                'Round robin does not accept knockout classification_places')
        local_entries = [e for e in entries if e['project_id'] == pid]
        require(len(local_entries) >= 2, 'Each project requires at least two competition slots')
        size_per_entry = p.get('entry_size')
        require(type(size_per_entry) is int and size_per_entry in (1, 2), 'entry_size must be 1 or 2; expand team ties to known submatches')
        known = [i for e in local_entries for i in e.get('member_ids', [])]
        require(len(known) == len(set(known)), 'Athlete entered twice in the same project')
        require(all(e.get('kind') == 'placeholder' or len(e['member_ids']) == size_per_entry for e in local_entries), 'Entry member count differs from project')
        d = draws.get(pid)
        require(d is not None, 'A confirmed draw is required for every project')
        require(d.get('format') == ('knockout' if p['format'] == 'knockout' else 'groups'), 'Draw format mismatch')
        assigned = deepcopy(d.get('assignments', []))
        require(len(assigned) == len(local_entries) and {a['entry_id'] for a in assigned} == {e['id'] for e in local_entries}, 'Draw must include every entry exactly once')
        require(all(type(a.get('group')) is int and a['group'] >= 1 and type(a.get('position')) is int and a['position'] >= 1 for a in assigned), 'Draw positions and groups must be positive integers')
        require(len({(a['group'], a['position']) for a in assigned}) == len(assigned), 'Duplicate draw slot')
        assigned.sort(key=lambda a: (a['group'], a['position']))
        topology = {k: p.get(k) for k in ('id', 'format', 'third_place', 'legs', 'advance_per_group', 'knockout_slots', 'enforce_round_order')}
        if 'classification_places' in p:
            topology['classification_places'] = p['classification_places']
        topology['assignments'] = [{k: a[k] for k in ('entry_id', 'group', 'position')} for a in assigned]
        topology['bracket_size'] = d.get('bracket_size')
        revision = digest(topology)[:16]
        p['topology_revision'] = revision
        p['platform_project_id'] = None
        p['entry_ids'] = [e['id'] for e in local_entries]
        out['projects'].append(p)
        by_match = {}

        def entrant(a):
            e = by_entry[a['entry_id']]
            return {'kind': 'entry', 'entry_id': e['id'], 'label': e['label'],
                    'club': e.get('club', ''), 'draw_position': a['position'],
                    'athletes': e.get('member_ids', []), 'possible_athletes': e.get('member_ids', []),
                    'possible_entry_ids': [e['id']], 'unresolved_identity_entries': [e['id']] if e.get('kind') == 'placeholder' else []}

        def match(a, b, stage, round_no, position, group_id=None, rank=None, placement=None, extra_predecessors=()):
            mid = stable_id('match', config['event_id'], pid, revision, stage, group_id, round_no, position, rank)
            label = f'{pid}-S{stage}-' + (f'G{group_id}-' if group_id else '') + f'R{round_no:02d}-M{position:02d}'
            sides = [deepcopy(a), deepcopy(b)]
            predecessors = set(extra_predecessors)
            for side in sides:
                if side['kind'] in ('winner', 'loser'):
                    predecessors.add(side['match_id'])
                elif side['kind'] == 'group_rank':
                    predecessors.update(side['match_ids'])
            m = {'id': mid, 'code': label, 'project_id': pid, 'project_name': p['name'],
                 'sport': p['sport'], 'group_name': p.get('group_name', p['name']),
                 'stage': stage, 'extra': 1, 'round': round_no, 'position': position,
                 'group_id': group_id, 'rank': rank, 'sides': sides,
                 'predecessors': sorted(predecessors),
                 'athletes': sorted({i for s in sides for i in s.get('athletes', [])}),
                 'possible_athletes': sorted({i for s in sides for i in s['possible_athletes']}),
                 'possible_entry_ids': sorted({i for s in sides for i in s['possible_entry_ids']}),
                 'unresolved_identity_entries': sorted({i for s in sides for i in s.get('unresolved_identity_entries', [])}),
                 'status': 'pending', 'platform_match_id': None}
            # Round-order constraints can carry competitors beyond this match's sides.
            for parent in extra_predecessors:
                m['possible_athletes'] = sorted(set(m['possible_athletes']) | set(by_match[parent]['possible_athletes']))
            if placement:
                m['placement'] = placement
            if 'duration_minutes' in p:
                require(type(p['duration_minutes']) is int and p['duration_minutes'] > 0, 'duration_minutes must be positive integer')
                m['duration_minutes'] = p['duration_minutes']
            out['matches'].append(m)
            by_match[mid] = m
            return {'kind': 'winner', 'match_id': mid, 'label': label + '胜者',
                    'athletes': [], 'possible_athletes': m['possible_athletes'],
                    'possible_entry_ids': m['possible_entry_ids'],
                    'unresolved_identity_entries': m['unresolved_identity_entries']}

        def knockout(slots, stage, bronze):
            size = len(slots)
            count = sum(s is not None for s in slots)
            require(size >= 2 and not size & (size - 1) and size == 1 << (count - 1).bit_length(), 'Use minimal power-of-two bracket size')
            require(all(slots[i] or slots[i + 1] for i in range(0, size, 2)), 'Double empty first-round pair; correct confirmed draw')
            require(type(bronze) is bool, 'third_place must be explicit boolean')
            require(not bronze or count >= 4, 'Third-place match requires at least four entrants')
            classification = p.get('classification_places')
            require(classification is None or (type(classification) is int and classification == 8), 'Explicit classification_places currently supports 8')
            require(classification is None or (bronze and count >= 8), 'Top-eight classification requires eight entrants and third_place')
            current, round_no, semifinals, quarterfinals = slots, 0, [], []
            while len(current) > 1:
                round_no += 1
                nxt = []
                for i in range(0, len(current), 2):
                    a, b = current[i:i + 2]
                    if not a or not b:
                        side = a or b
                        out['automatic_advances'].append({'project_id': pid, 'stage': stage,
                            'round': round_no, 'position': i // 2 + 1, 'side': deepcopy(side),
                            'reason': 'bye', 'counts_as_played_match': False})
                        nxt.append(side)
                        continue
                    winner = match(a, b, stage, round_no, i // 2 + 1, rank=f'1-{len(current)}',
                                   placement={'winner': 1, 'loser': 2} if len(current) == 2 else None)
                    if len(current) == 4:
                        semifinals.append(winner)
                    if len(current) == 8:
                        quarterfinals.append(winner)
                    nxt.append(winner)
                current = nxt
            if bronze:
                require(len(semifinals) == 2, 'Third-place match requires two played semifinals')
                losing = [{**s, 'kind': 'loser', 'label': s['label'][:-2] + '负者'} for s in semifinals]
                third = match(*losing, stage, round_no, 2, rank='3-4', placement={'winner': 3, 'loser': 4})
                # Conservative candidate pools may overlap for qualifiers from one group.
                pools = [set(s['possible_athletes']) for s in semifinals]
                if pools[0] and pools[1] and not pools[0] & pools[1]:
                    out['outcome_disjoint_pairs'].append([current[0]['match_id'], third['match_id']])
            if classification == 8:
                require(len(quarterfinals) == 4, 'Top-eight classification requires four played quarterfinals')
                losing = [{**s, 'kind': 'loser', 'label': s['label'][:-2] + '负者'} for s in quarterfinals]
                five_semis = [match(losing[i], losing[i + 1], stage, round_no - 1, i // 2 + 3, rank='5-8')
                              for i in (0, 2)]
                fifth = match(*five_semis, stage, round_no, 3, rank='5-6', placement={'winner': 5, 'loser': 6})
                five_losers = [{**s, 'kind': 'loser', 'label': s['label'][:-2] + '负者'} for s in five_semis]
                seventh = match(*five_losers, stage, round_no, 4, rank='7-8', placement={'winner': 7, 'loser': 8})
                pools = [set(s['possible_athletes']) for s in five_semis]
                if pools[0] and pools[1] and not pools[0] & pools[1]:
                    out['outcome_disjoint_pairs'].append([fifth['match_id'], seventh['match_id']])

        if p['format'] == 'knockout':
            require(all(a['group'] == 1 for a in assigned), 'Knockout draw uses group 1')
            size = d.get('bracket_size')
            require(type(size) is int and size >= len(assigned), 'Confirmed bracket_size required')
            require(all(a['position'] <= size for a in assigned), 'Draw position exceeds bracket')
            positions = {a['position']: entrant(a) for a in assigned}
            actual_byes = [n for n in range(1, size + 1) if n not in positions]
            require('byes' not in d or d['byes'] == actual_byes, 'Draw bye list inconsistent with occupied positions')
            knockout([positions.get(i) for i in range(1, size + 1)], 1, p.get('third_place'))
            continue

        group_numbers = sorted({a['group'] for a in assigned})
        require(group_numbers == list(range(1, len(group_numbers) + 1)), 'Group numbers must be contiguous')
        require(p['format'] != 'round_robin' or len(group_numbers) == 1, 'Round robin uses one group; use groups_knockout for multiple groups')
        legs = p.get('legs', 1)
        require(type(legs) is int and legs in (1, 2), 'Round robin supports one or two legs')
        groups = {}
        for g in group_numbers:
            assignments = [a for a in assigned if a['group'] == g]
            require([a['position'] for a in assignments] == list(range(1, len(assignments) + 1)), 'Group positions must be contiguous')
            require(len(assignments) >= 2, 'Each round robin group needs at least two entrants')
            gid = f'{pid}:G{g}'
            group = {'id': gid, 'number': g, 'project_id': pid, 'entry_ids': [a['entry_id'] for a in assignments],
                     'match_ids': [], 'expected_meetings': legs, 'ranking': deepcopy(p.get('ranking', {}))}
            groups[g] = group
            out['groups'].append(group)
            rotation = [entrant(a) for a in assignments]
            if len(rotation) % 2:
                rotation.append(None)
            first_leg, previous_round = [], []
            for r in range(1, len(rotation)):
                pairs = [(rotation[i], rotation[-i - 1]) for i in range(len(rotation) // 2)]
                first_leg.append(pairs)
                rotation = [rotation[0], rotation[-1], *rotation[1:-1]]
            for leg in range(legs):
                for offset, pairs in enumerate(first_leg, 1):
                    current_round = []
                    for pos, pair in enumerate(pairs, 1):
                        a, b = pair if leg == 0 else pair[::-1]
                        if a is None or b is None:
                            continue
                        winner = match(a, b, 1, leg * len(first_leg) + offset, pos, gid,
                                       extra_predecessors=previous_round if p.get('enforce_round_order') else ())
                        current_round.append(winner['match_id'])
                    group['match_ids'].extend(current_round)
                    previous_round = current_round
        if p['format'] == 'round_robin':
            continue
        advance = p.get('advance_per_group')
        require(type(advance) is int and advance >= 1 and all(advance <= len(g['entry_ids']) for g in groups.values()), 'Confirm valid advance_per_group')
        crossover = p.get('knockout_slots')
        require(isinstance(crossover, list), 'Explicit knockout_slots required; no implicit crossover or redraw')
        expected = {(g, rank) for g in groups for rank in range(1, advance + 1)}
        refs = [(s.get('group'), s.get('rank')) for s in crossover if s is not None]
        require(len(refs) == len(expected) and set(refs) == expected, 'Each qualifying group rank must appear exactly once')
        slots = []
        for s in crossover:
            if s is None:
                slots.append(None)
                continue
            group = groups[s['group']]
            pool = [by_entry[e] for e in group['entry_ids']]
            slots.append({'kind': 'group_rank', 'group_id': group['id'], 'rank': s['rank'],
                          'match_ids': group['match_ids'], 'label': f"{p['name']}第{s['group']}组第{s['rank']}名",
                          'athletes': [], 'possible_athletes': sorted({m for e in pool for m in e.get('member_ids', [])}),
                          'possible_entry_ids': group['entry_ids'],
                          'unresolved_identity_entries': [e['id'] for e in pool if e.get('kind') == 'placeholder']})
        knockout(slots, 2, p.get('third_place'))
    out['validation'] = validate(out)
    return out


def validate(data):
    """Check serialized graph and member identity invariants without engine state."""
    matches = data['matches']
    ids = {m['id'] for m in matches}
    require(len(ids) == len(matches), 'Duplicate match ID')
    by_id = {m['id']: m for m in matches}
    entries = {e['id']: e for e in data['entries']}
    groups = {g['id']: g for g in data['groups']}
    visiting, done = set(), set()

    def visit(mid):
        require(mid not in visiting, 'Cyclic dependencies')
        if mid in done:
            return
        visiting.add(mid)
        m = by_id[mid]
        require(len(m['sides']) == 2, 'Each played match needs two sides')
        require(set(m['athletes']) <= set(m['possible_athletes']), 'Known athletes absent from candidates')
        known_list = [a for s in m['sides'] for a in s.get('athletes', [])]
        require(len(known_list) == len(set(known_list)), 'Athlete faces self')
        require(set(known_list) == set(m['athletes']), 'Known athlete union mismatch')
        expected_parents = set()
        for side in m['sides']:
            kind = side['kind']
            if kind == 'entry':
                require(side['entry_id'] in entries, 'Unknown entry source')
                e = entries[side['entry_id']]
                require(e['project_id'] == m['project_id'] and set(side['athletes']) == set(e.get('member_ids', [])), 'Entry member/project mismatch')
            elif kind in ('winner', 'loser'):
                parent = side['match_id']
                require(parent in by_id and by_id[parent]['project_id'] == m['project_id'], 'Unknown or cross-project source')
                require(by_id[parent]['stage'] < m['stage'] or (by_id[parent]['stage'] == m['stage'] and by_id[parent]['round'] < m['round']), 'Knockout source must be earlier round/stage')
                expected_parents.add(parent)
            elif kind == 'group_rank':
                require(side['group_id'] in groups, 'Unknown group source')
                group = groups[side['group_id']]
                require(group['project_id'] == m['project_id'] and set(side['match_ids']) == set(group['match_ids']), 'Qualification dependency incomplete')
                require(1 <= side['rank'] <= len(group['entry_ids']), 'Qualification rank out of range')
                expected_parents.update(side['match_ids'])
            else:
                raise ValueError('Unknown side kind: ' + str(kind))
        require(expected_parents <= set(m['predecessors']), 'Missing source dependency')
        for parent in m['predecessors']:
            require(parent in ids, 'Missing predecessor')
            visit(parent)
            require(set(by_id[parent]['possible_athletes']) <= set(m['possible_athletes']), 'Predecessor candidates lost')
        visiting.remove(mid)
        done.add(mid)
    for mid in ids:
        visit(mid)
    return {'passed': True, 'matches': len(matches), 'automatic_advances': len(data['automatic_advances']),
            'all_platform_ids_empty': all(m.get('platform_match_id') is None for m in matches),
            'unresolved_identities': len(data.get('unresolved_identity_entries', [])),
            'scope': 'graph and known identities; scoring and qualification ranking evaluated separately'}


def export_excel(data, path):
    """Review workbook with the familiar 17 fields, never a platform upload file."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    validate(data)
    path = Path(path)
    require(not path.exists(), 'Output exists; choose a new revision filename')
    book = Workbook()
    book.remove(book.active)

    def sheet(title, headers, rows):
        ws = book.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for row in ws:
            for c in row:
                c.font = Font(name='Arial', size=10)
                c.alignment = Alignment(vertical='center', wrap_text=True)
                if isinstance(c.value, str):
                    c.data_type = 's'
        for c in ws[1]:
            c.font = Font(name='Arial', bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='215E75')
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 25 if col > 1 else 10
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = '1:1'
        return ws
    rows = []
    group_numbers = {g['id']: g['number'] for g in data['groups']}
    for n, m in enumerate(data['matches'], 1):
        a, b = m['sides']
        rows.append([n, data['event_name'], None, m['group_name'], m['project_name'], group_numbers.get(m.get('group_id'), '-'),
                     m['round'], m['stage'], m['extra'], m.get('rank'),
                     a.get('club', ''), a['label'], b.get('club', ''), b['label'], None, None, m['code']])
    sheet('赛程审阅17列', HEADERS17, rows)
    sheet('本地场次与依赖', ['赛号', '本地场次ID', '本地项目ID', '前置场次ID', '已确定运动员ID', '候选运动员ID', '未知身份预留项'],
          [[m['code'], m['id'], m['project_id'], '\n'.join(m['predecessors']), '\n'.join(m['athletes']),
            '\n'.join(m['possible_athletes']), '\n'.join(m['unresolved_identity_entries'])] for m in data['matches']])
    sheet('轮空自动晋级', ['项目', '阶段', '轮次', '位置', '晋级项', '计入已赛场数'],
          [[b['project_id'], b['stage'], b['round'], b['position'], b['side']['label'], '否'] for b in data['automatic_advances']])
    sheet('使用说明', ['事项', '说明'], [
        ['性质', '完全离线赛程审阅表；字段沿用默认模板风格，平台项目ID与场次ID为空。'],
        ['对阵', '种子、单位回避与签位由已确认抽签决定；本引擎不重新抽签。'],
        ['时间场地', '本文件是对阵图；另交 schedule_engine 按休息、时段、场地编排。'],
        ['预留', '未知身份不能保证跨项目无冲突；替换后重新核对全赛事身份、休息和回避。'],
        ['循环晋级', '全部组内场次结束并按确认的排名规则解出名次后才能确定晋级人选。'],
        ['输入摘要', data['input_sha256']]])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    check = load_workbook(path, data_only=False)
    require(check.worksheets[0].max_row == len(rows) + 1, 'Saved workbook lost rows')
    require(all(check.worksheets[0].cell(r, c).value is None for r in range(2, len(rows) + 2) for c in (3, 15, 16)), 'Platform ID leaked into local workbook')
    require([check.worksheets[0].cell(r, 17).value for r in range(2, len(rows) + 2)] == [m['code'] for m in data['matches']], 'Saved match codes differ')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--xlsx', type=Path)
    args = parser.parse_args()
    require(not args.output.exists() and (not args.xlsx or not args.xlsx.exists()), 'Choose unused output filenames')
    data = generate(json.loads(args.input.read_text(encoding='utf-8')))
    if args.xlsx:
        export_excel(data, args.xlsx)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(data['validation'], ensure_ascii=False))


if __name__ == '__main__':
    main()
